#!/usr/bin/env python3
"""
PostgreSQL-only async memory CLI for Codex, Claude Code, and local agents.

Design constraints:
- raw asyncpg only
- no ORM or validation-model database layer
- no SQLite/Chroma side database
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import json
import math
import mimetypes
import os
import re
import shlex
import shutil
import subprocess
import sys
import tempfile
import wave
import zipfile
import zlib
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import asyncpg


BACKEND_DIR = Path(__file__).resolve().parent
LOCAL_ENV_PREFIXES = ("MEMORY_", "AGENT_MEMORY_")


def load_local_env(path: Path) -> None:
    """Load MEMORY_* and AGENT_MEMORY_* settings from a gitignored .env file.

    Real environment variables always win. Only the file next to this script
    (or AGENT_MEMORY_ENV_FILE) is read, never the hook's cwd, so a project's own
    .env cannot redirect memory writes.
    """
    try:
        lines = path.read_text(encoding="utf-8-sig").splitlines()
    except OSError:
        return
    for line in lines:
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.removeprefix("export ").strip()
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
            value = value[1:-1]
        if key.startswith(LOCAL_ENV_PREFIXES):
            os.environ.setdefault(key, value)


load_local_env(Path(os.getenv("AGENT_MEMORY_ENV_FILE") or BACKEND_DIR / ".env"))

VECTOR_DIMS = 384
DEFAULT_SEARCH_CANDIDATES = 400
MAX_SEARCH_CANDIDATES = 2000
SEARCH_CACHE_VERSION = "hybrid-v2-strict"
DEFAULT_MIN_SEARCH_SCORE = float(os.getenv("AGENT_MEMORY_MIN_SEARCH_SCORE", "0.18"))
MAX_INDEXED_CONTENT_CHARS = max(10000, int(os.getenv("AGENT_MEMORY_MAX_INDEXED_CONTENT_CHARS", "180000")))
DEFAULT_SPACE = os.getenv("AGENT_MEMORY_SPACE", "global").strip() or "global"
DEFAULT_DB = {
    "host": "127.0.0.1",
    "port": 25490,
    "database": "agent_memory",
    "user": "agent_memory",
}
LOG_DIR = Path(os.getenv("AGENT_MEMORY_LOG_DIR", str(BACKEND_DIR / "logs")))


def json_default(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def read_password_file(path_text: str) -> str:
    if not path_text:
        return ""
    try:
        return Path(path_text).expanduser().read_text(encoding="utf-8-sig").strip()
    except OSError:
        return ""


def read_password_command(command: str) -> str:
    """Run a credential helper that prints the database password on stdout.

    `{python}` as the first word is replaced with the running interpreter, so a
    Python vault CLI runs under the same environment as the hook.
    """
    if not command:
        return ""
    try:
        argv = shlex.split(command, posix=True)
    except ValueError:
        return ""
    if not argv:
        return ""
    if argv[0] == "{python}":
        argv[0] = sys.executable
    try:
        result = subprocess.run(argv, capture_output=True, text=True, timeout=10, check=False)
    except (OSError, subprocess.SubprocessError):
        return ""
    return result.stdout.strip() if result.returncode == 0 else ""


def db_config() -> dict[str, Any]:
    password = (
        os.getenv("MEMORY_DB_PASSWORD")
        or os.getenv("PGPASSWORD")
        or read_password_file(os.getenv("MEMORY_DB_PASSWORD_FILE", ""))
        or read_password_command(os.getenv("MEMORY_DB_PASSWORD_COMMAND", ""))
    )
    return {
        "host": os.getenv("MEMORY_DB_HOST", DEFAULT_DB["host"]),
        "port": int(os.getenv("MEMORY_DB_PORT", str(DEFAULT_DB["port"]))),
        "database": os.getenv("MEMORY_DB_NAME", DEFAULT_DB["database"]),
        "user": os.getenv("MEMORY_DB_USER", DEFAULT_DB["user"]),
        "password": password,
        "timeout": float(os.getenv("MEMORY_DB_CONNECT_TIMEOUT", "10")),
    }


def safe_argv(argv: list[str]) -> list[str]:
    safe: list[str] = []
    for index, value in enumerate(argv):
        text = str(value)
        if index < 2:
            safe.append(text)
        elif text.startswith("--"):
            safe.append(text.split("=", 1)[0] + ("=<redacted-arg>" if "=" in text else ""))
        elif text.startswith("-") and len(text) <= 8:
            safe.append(text)
        else:
            digest = hashlib.sha256(text.encode("utf-8", errors="backslashreplace")).hexdigest()
            safe.append(f"<arg chars={len(text)} sha256={digest}>")
    return safe


def append_jsonl_locked(path: Path, record: dict[str, Any]) -> None:
    line = json.dumps(record, ensure_ascii=False, default=json_default, sort_keys=True) + "\n"
    with path.open("a+", encoding="utf-8", newline="\n") as handle:
        if os.name == "nt":
            import msvcrt

            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_LOCK, 1)
        else:
            import fcntl

            fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        try:
            handle.seek(0, os.SEEK_END)
            handle.write(line)
            handle.flush()
            os.fsync(handle.fileno())
        finally:
            if os.name == "nt":
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def write_error_log(exc: Exception) -> None:
    record = {
        "created_at": datetime.now(UTC).isoformat(),
        "error": str(exc),
        "argv": safe_argv(sys.argv),
    }
    for directory in (LOG_DIR, Path("/tmp/agent-memory/logs"), Path.cwd() / "logs"):
        try:
            directory.mkdir(parents=True, exist_ok=True)
            append_jsonl_locked(directory / "agent_memory_errors.jsonl", record)
            return
        except Exception:
            continue


async def connect() -> asyncpg.Connection:
    conn = await asyncpg.connect(**db_config())
    await conn.set_type_codec(
        "jsonb",
        encoder=lambda value: json.dumps(value, ensure_ascii=False, default=json_default),
        decoder=json.loads,
        schema="pg_catalog",
    )
    await conn.set_type_codec(
        "json",
        encoder=lambda value: json.dumps(value, ensure_ascii=False, default=json_default),
        decoder=json.loads,
        schema="pg_catalog",
    )
    return conn


TOKEN_RE = re.compile(r"[a-z0-9][a-z0-9_./:-]*", re.IGNORECASE)
ENTITY_RE = re.compile(
    r"\b(?:[A-Z][A-Za-z0-9_.:/-]{2,}|[a-z0-9]+(?:[-_/.:][a-z0-9]+)+)"
    r"(?:\s+(?:[A-Z][A-Za-z0-9_.:/-]{2,}|[a-z0-9]+(?:[-_/.:][a-z0-9]+)+)){0,4}\b"
)


def stem_token(token: str) -> list[str]:
    token = token.lower()
    variants = {token}
    for suffix in ("ingly", "edly", "ing", "ed", "ies", "s"):
        if len(token) > len(suffix) + 3 and token.endswith(suffix):
            if suffix == "ies":
                variants.add(token[: -len(suffix)] + "y")
            else:
                variants.add(token[: -len(suffix)])
    return list(variants)


def stable_hash(text: str, salt: str = "") -> int:
    digest = hashlib.blake2b((salt + text).encode("utf-8"), digest_size=8).digest()
    return int.from_bytes(digest, "big", signed=False)


def add_feature(vec: list[float], feature: str, weight: float) -> None:
    h = stable_hash(feature)
    idx = h % VECTOR_DIMS
    sign = 1.0 if ((h >> 8) & 1) else -1.0
    vec[idx] += sign * weight


def cpu_embedding(text: str) -> list[float]:
    """Dependency-free CPU feature-hash embedding for local pgvector search."""
    normalized = text.lower()
    tokens: list[str] = []
    for raw in TOKEN_RE.findall(normalized):
        if len(raw) > 96:
            raw = raw[:96]
        tokens.extend(stem_token(raw))

    vec = [0.0] * VECTOR_DIMS
    for token in tokens:
        add_feature(vec, "tok:" + token, 1.0)
        if len(token) >= 5:
            for i in range(len(token) - 2):
                add_feature(vec, "tri:" + token[i : i + 3], 0.35)

    for left, right in zip(tokens, tokens[1:]):
        add_feature(vec, "bi:" + left + " " + right, 0.8)

    norm = math.sqrt(sum(v * v for v in vec))
    if norm == 0:
        return vec
    return [round(v / norm, 8) for v in vec]


def vector_literal(values: list[float]) -> str:
    return "[" + ",".join(f"{v:.8f}" for v in values) + "]"


def search_candidate_limit(limit: int) -> int:
    """Bound the candidate pool used before final blended scoring."""
    try:
        configured = int(os.getenv("AGENT_MEMORY_SEARCH_CANDIDATES", str(DEFAULT_SEARCH_CANDIDATES)))
    except ValueError:
        configured = DEFAULT_SEARCH_CANDIDATES
    try:
        multiplier = int(os.getenv("AGENT_MEMORY_SEARCH_CANDIDATE_MULTIPLIER", "40"))
    except ValueError:
        multiplier = 40
    try:
        maximum = int(os.getenv("AGENT_MEMORY_SEARCH_CANDIDATE_MAX", str(MAX_SEARCH_CANDIDATES)))
    except ValueError:
        maximum = MAX_SEARCH_CANDIDATES
    minimum = max(limit, limit * max(multiplier, 1), configured)
    return max(limit, min(maximum, minimum))


def search_cache_enabled() -> bool:
    return os.getenv("AGENT_MEMORY_SEARCH_CACHE", "1").strip().lower() not in {"0", "false", "no", "off"}


def search_cache_key(
    query: str,
    *,
    limit: int,
    space: str,
    all_spaces: bool,
    agent: str,
    event_type: str,
    memory_kind: str,
    tags: list[str] | None,
    min_score: float,
    candidate_limit: int,
    phrase_candidates: bool,
) -> str:
    cache_spec = {
        "version": SEARCH_CACHE_VERSION,
        "query": query,
        "limit": limit,
        "space": "" if all_spaces else normalize_space(space),
        "all_spaces": all_spaces,
        "agent": agent,
        "event_type": event_type,
        "memory_kind": memory_kind,
        "tags": sorted(set(tags or [])),
        "min_score": min_score,
        "candidate_limit": candidate_limit,
        "phrase_candidates": phrase_candidates,
    }
    encoded = json.dumps(cache_spec, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8", errors="backslashreplace")).hexdigest()


def restore_cached_search_results(payload: Any) -> list[dict[str, Any]]:
    if not isinstance(payload, list):
        return []
    results: list[dict[str, Any]] = []
    for raw_item in payload:
        if not isinstance(raw_item, dict):
            continue
        item = dict(raw_item)
        created_at = item.get("created_at")
        if isinstance(created_at, str):
            try:
                item["created_at"] = datetime.fromisoformat(created_at.replace("Z", "+00:00"))
            except ValueError:
                pass
        results.append(item)
    return results


def summarize(content: str, limit: int = 280) -> str:
    one_line = re.sub(r"\s+", " ", content).strip()
    if len(one_line) <= limit:
        return one_line
    return one_line[: limit - 3] + "..."


def compact_indexed_content(content: str) -> tuple[str, dict[str, Any] | None]:
    if len(content) <= MAX_INDEXED_CONTENT_CHARS:
        return content, None
    digest = hashlib.sha256(content.encode("utf-8", errors="backslashreplace")).hexdigest()
    tail_chars = min(8000, MAX_INDEXED_CONTENT_CHARS // 10)
    head_chars = MAX_INDEXED_CONTENT_CHARS - tail_chars - 180
    marker = (
        f"\n<agent-memory full content archived: original_chars={len(content)} "
        f"sha256={digest}>\n"
    )
    compact = content[:head_chars] + marker + content[-tail_chars:]
    return compact, {
        "payload_kind": "oversize_event_content",
        "payload_sha256": digest,
        "original_chars": len(content),
        "indexed_chars": len(compact),
        "stored": True,
    }


def normalize_space(space: str | None) -> str:
    value = (space or DEFAULT_SPACE or "global").strip().lower()
    value = re.sub(r"[^a-z0-9_.:/-]+", "-", value).strip("-")
    return value[:160] or "global"


def normalize_entity_name(name: str) -> str:
    return re.sub(r"\s+", " ", (name or "").strip())[:240]


def extract_entities(content: str, tags: list[str] | None = None, metadata: dict[str, Any] | None = None) -> list[dict[str, str]]:
    """Extract lightweight concept candidates for the memory knowledge graph."""
    candidates: dict[str, dict[str, str]] = {}
    for item in (metadata or {}).get("entities", []) if isinstance(metadata, dict) else []:
        if isinstance(item, dict):
            name = normalize_entity_name(str(item.get("name") or ""))
            kind = str(item.get("kind") or "concept").strip().lower() or "concept"
        else:
            name = normalize_entity_name(str(item))
            kind = "concept"
        if name:
            candidates[name.lower()] = {"name": name, "kind": kind[:80]}

    for tag in tags or []:
        tag_text = normalize_entity_name(str(tag))
        if tag_text and len(tag_text) >= 4 and tag_text not in {"asyncpg", "memory"}:
            candidates.setdefault(tag_text.lower(), {"name": tag_text, "kind": "tag"})

    for match in ENTITY_RE.findall(content or ""):
        name = normalize_entity_name(match)
        lowered = name.lower()
        if len(name) < 4 or lowered in {"http", "https", "json", "true", "false", "null"}:
            continue
        candidates.setdefault(lowered, {"name": name, "kind": "concept"})
        if len(candidates) >= 32:
            break
    return list(candidates.values())[:32]


async def _execute_ddl(conn: asyncpg.Connection, sql: str) -> None:
    try:
        await conn.execute(sql)
    except (
        asyncpg.DuplicateObjectError,
        asyncpg.DuplicateTableError,
        asyncpg.UniqueViolationError,
        asyncpg.ObjectInUseError,
    ):
        return


async def init_db_async() -> dict[str, Any]:
    conn = await connect()
    try:
        await _execute_ddl(conn, "CREATE EXTENSION IF NOT EXISTS vector")
        await _execute_ddl(conn, "CREATE EXTENSION IF NOT EXISTS pg_trgm")
        await _execute_ddl(
            conn,
            """
            CREATE TABLE IF NOT EXISTS agent_memory_spaces (
                space TEXT PRIMARY KEY,
                description TEXT,
                parent_space TEXT REFERENCES agent_memory_spaces(space) ON DELETE SET NULL,
                metadata JSONB NOT NULL DEFAULT '{}'::JSONB,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """,
        )
        await conn.execute(
            """
            INSERT INTO agent_memory_spaces (space, description)
            VALUES ('global', 'Default shared memory space for unscoped legacy events.')
            ON CONFLICT (space) DO NOTHING
            """
        )
        await _execute_ddl(
            conn,
            """
            CREATE TABLE IF NOT EXISTS agent_memory_events (
                id BIGSERIAL PRIMARY KEY,
                event_uid TEXT UNIQUE NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                space TEXT NOT NULL DEFAULT 'global',
                session_id TEXT,
                agent TEXT NOT NULL,
                event_type TEXT NOT NULL,
                memory_kind TEXT NOT NULL DEFAULT 'event',
                role TEXT,
                source TEXT,
                content TEXT NOT NULL,
                summary TEXT,
                tags TEXT[] NOT NULL DEFAULT ARRAY[]::TEXT[],
                metadata JSONB NOT NULL DEFAULT '{}'::JSONB,
                embedding VECTOR(384),
                search_vector TSVECTOR GENERATED ALWAYS AS (
                    to_tsvector('english', coalesce(content, '') || ' ' || coalesce(summary, ''))
                ) STORED
            )
            """,
        )
        await _execute_ddl(conn, "ALTER TABLE agent_memory_events ADD COLUMN IF NOT EXISTS space TEXT NOT NULL DEFAULT 'global'")
        await _execute_ddl(conn, "ALTER TABLE agent_memory_events ADD COLUMN IF NOT EXISTS memory_kind TEXT NOT NULL DEFAULT 'event'")
        await _execute_ddl(
            conn,
            """
            CREATE TABLE IF NOT EXISTS agent_memory_search_state (
                singleton BOOLEAN PRIMARY KEY DEFAULT TRUE CHECK (singleton),
                revision BIGINT NOT NULL DEFAULT 0,
                mutation_revision BIGINT NOT NULL DEFAULT 0,
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """,
        )
        await _execute_ddl(
            conn,
            "ALTER TABLE agent_memory_search_state ADD COLUMN IF NOT EXISTS mutation_revision BIGINT NOT NULL DEFAULT 0",
        )
        await conn.execute(
            """
            INSERT INTO agent_memory_search_state (singleton, revision)
            VALUES (TRUE, 0)
            ON CONFLICT (singleton) DO NOTHING
            """
        )
        await _execute_ddl(
            conn,
            """
            CREATE TABLE IF NOT EXISTS agent_memory_search_cache (
                cache_key TEXT PRIMARY KEY,
                revision BIGINT NOT NULL,
                mutation_revision BIGINT NOT NULL DEFAULT 0,
                max_event_id BIGINT NOT NULL DEFAULT 0,
                results JSONB NOT NULL,
                result_count INTEGER NOT NULL,
                hit_count BIGINT NOT NULL DEFAULT 0,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                refreshed_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                last_accessed_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """,
        )
        await _execute_ddl(
            conn,
            "ALTER TABLE agent_memory_search_cache ADD COLUMN IF NOT EXISTS mutation_revision BIGINT NOT NULL DEFAULT 0",
        )
        await _execute_ddl(
            conn,
            "ALTER TABLE agent_memory_search_cache ADD COLUMN IF NOT EXISTS max_event_id BIGINT NOT NULL DEFAULT 0",
        )
        await conn.execute(
            """
            CREATE OR REPLACE FUNCTION agent_memory_bump_search_revision()
            RETURNS TRIGGER
            LANGUAGE plpgsql
            AS $$
            BEGIN
                UPDATE agent_memory_search_state
                SET revision = revision + 1, updated_at = NOW()
                WHERE singleton = TRUE;
                RETURN NULL;
            END;
            $$
            """
        )
        await conn.execute(
            """
            CREATE OR REPLACE FUNCTION agent_memory_bump_search_mutation_revision()
            RETURNS TRIGGER
            LANGUAGE plpgsql
            AS $$
            BEGIN
                UPDATE agent_memory_search_state
                SET mutation_revision = mutation_revision + 1, updated_at = NOW()
                WHERE singleton = TRUE;
                RETURN NULL;
            END;
            $$
            """
        )
        await conn.execute(
            """
            DO $$
            BEGIN
                IF NOT EXISTS (
                    SELECT 1
                    FROM pg_trigger
                    WHERE tgname = 'agent_memory_events_search_revision_trg'
                      AND tgrelid = 'agent_memory_events'::regclass
                      AND NOT tgisinternal
                ) THEN
                    CREATE TRIGGER agent_memory_events_search_revision_trg
                    AFTER INSERT OR UPDATE OR DELETE ON agent_memory_events
                    FOR EACH STATEMENT
                    EXECUTE FUNCTION agent_memory_bump_search_revision();
                END IF;
            END;
            $$
            """
        )
        await conn.execute(
            """
            DO $$
            BEGIN
                IF NOT EXISTS (
                    SELECT 1
                    FROM pg_trigger
                    WHERE tgname = 'agent_memory_events_search_mutation_trg'
                      AND tgrelid = 'agent_memory_events'::regclass
                      AND NOT tgisinternal
                ) THEN
                    CREATE TRIGGER agent_memory_events_search_mutation_trg
                    AFTER UPDATE OR DELETE ON agent_memory_events
                    FOR EACH ROW
                    EXECUTE FUNCTION agent_memory_bump_search_mutation_revision();
                END IF;
            END;
            $$
            """
        )
        await _execute_ddl(
            conn,
            """
            CREATE TABLE IF NOT EXISTS agent_memory_entities (
                id BIGSERIAL PRIMARY KEY,
                space TEXT NOT NULL DEFAULT 'global',
                name TEXT NOT NULL,
                kind TEXT NOT NULL DEFAULT 'concept',
                summary TEXT,
                metadata JSONB NOT NULL DEFAULT '{}'::JSONB,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """,
        )
        await _execute_ddl(
            conn,
            """
            CREATE TABLE IF NOT EXISTS agent_memory_entity_mentions (
                event_id BIGINT NOT NULL REFERENCES agent_memory_events(id) ON DELETE CASCADE,
                entity_id BIGINT NOT NULL REFERENCES agent_memory_entities(id) ON DELETE CASCADE,
                weight REAL NOT NULL DEFAULT 1.0,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                PRIMARY KEY (event_id, entity_id)
            )
            """,
        )
        await _execute_ddl(
            conn,
            """
            CREATE TABLE IF NOT EXISTS agent_memory_entity_edges (
                id BIGSERIAL PRIMARY KEY,
                space TEXT NOT NULL DEFAULT 'global',
                source_entity_id BIGINT NOT NULL REFERENCES agent_memory_entities(id) ON DELETE CASCADE,
                target_entity_id BIGINT NOT NULL REFERENCES agent_memory_entities(id) ON DELETE CASCADE,
                relation TEXT NOT NULL DEFAULT 'related_to',
                description TEXT,
                metadata JSONB NOT NULL DEFAULT '{}'::JSONB,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """,
        )
        await _execute_ddl(
            conn,
            """
            CREATE TABLE IF NOT EXISTS agent_memory_event_links (
                id BIGSERIAL PRIMARY KEY,
                source_event_id BIGINT NOT NULL REFERENCES agent_memory_events(id) ON DELETE CASCADE,
                target_event_id BIGINT NOT NULL REFERENCES agent_memory_events(id) ON DELETE CASCADE,
                link_type TEXT NOT NULL DEFAULT 'related_to',
                description TEXT,
                metadata JSONB NOT NULL DEFAULT '{}'::JSONB,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """,
        )
        await _execute_ddl(
            conn,
            """
            CREATE TABLE IF NOT EXISTS agent_memory_payload_archives (
                id BIGSERIAL PRIMARY KEY,
                event_id BIGINT REFERENCES agent_memory_events(id) ON DELETE CASCADE,
                event_uid TEXT NOT NULL,
                payload_sha256 TEXT NOT NULL,
                payload_kind TEXT NOT NULL DEFAULT 'hook_payload',
                payload_format TEXT NOT NULL DEFAULT 'json',
                compression TEXT NOT NULL DEFAULT 'zlib',
                raw_chars BIGINT NOT NULL,
                compressed_bytes BIGINT NOT NULL,
                redacted BOOLEAN NOT NULL DEFAULT TRUE,
                space TEXT NOT NULL DEFAULT 'global',
                session_id TEXT,
                agent TEXT NOT NULL,
                event_type TEXT NOT NULL,
                source TEXT,
                payload BYTEA NOT NULL,
                metadata JSONB NOT NULL DEFAULT '{}'::JSONB,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                UNIQUE (event_uid, payload_kind, payload_sha256)
            )
            """,
        )
        ddl = [
            """
            CREATE INDEX IF NOT EXISTS agent_memory_events_space_created_idx
            ON agent_memory_events (space, created_at DESC)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_events_space_kind_idx
            ON agent_memory_events (space, memory_kind, created_at DESC)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_events_created_idx
            ON agent_memory_events (created_at DESC)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_events_agent_type_idx
            ON agent_memory_events (agent, event_type, created_at DESC)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_events_tags_idx
            ON agent_memory_events USING GIN (tags)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_events_search_idx
            ON agent_memory_events USING GIN (search_vector)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_events_content_trgm_idx
            ON agent_memory_events USING GIN (content gin_trgm_ops)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_events_embedding_idx
            ON agent_memory_events USING hnsw (embedding vector_cosine_ops)
            WITH (m = 16, ef_construction = 100)
            """,
            """
            CREATE UNIQUE INDEX IF NOT EXISTS agent_memory_entities_space_name_kind_uidx
            ON agent_memory_entities (space, lower(name), kind)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_entities_space_idx
            ON agent_memory_entities (space, kind, updated_at DESC)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_entities_name_trgm_idx
            ON agent_memory_entities USING GIN (name gin_trgm_ops)
            """,
            """
            CREATE UNIQUE INDEX IF NOT EXISTS agent_memory_entity_edges_uidx
            ON agent_memory_entity_edges (space, source_entity_id, target_entity_id, relation)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_entity_edges_source_idx
            ON agent_memory_entity_edges (source_entity_id, relation)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_event_links_source_idx
            ON agent_memory_event_links (source_event_id, link_type)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_payload_archives_event_idx
            ON agent_memory_payload_archives (event_id, payload_kind, created_at DESC)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_payload_archives_hash_idx
            ON agent_memory_payload_archives (payload_sha256)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_payload_archives_space_idx
            ON agent_memory_payload_archives (space, created_at DESC)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_payload_archives_session_idx
            ON agent_memory_payload_archives (session_id, created_at DESC)
            """,
            """
            CREATE INDEX IF NOT EXISTS agent_memory_search_cache_accessed_idx
            ON agent_memory_search_cache (last_accessed_at DESC)
            """,
        ]
        for sql in ddl:
            await _execute_ddl(conn, sql)
        await _execute_ddl(conn, """
            CREATE INDEX CONCURRENTLY IF NOT EXISTS agent_memory_events_identity_gin_idx
            ON agent_memory_events USING GIN ((metadata->'identity'))
            WHERE metadata ? 'identity'
        """)
        return {"ok": True, "message": "agent_memory_events schema ready", "driver": "asyncpg"}
    finally:
        await conn.close()


def make_uid(
    agent: str,
    event_type: str,
    session_id: str,
    content: str,
    metadata: dict[str, Any],
    space: str = "",
    memory_kind: str = "event",
) -> str:
    seed = json.dumps(
        {
            "space": normalize_space(space),
            "agent": agent,
            "event_type": event_type,
            "memory_kind": memory_kind,
            "session_id": session_id,
            "content": content,
            "metadata": metadata,
        },
        sort_keys=True,
        ensure_ascii=False,
        default=json_default,
    )
    return hashlib.sha256(seed.encode("utf-8")).hexdigest()


async def ensure_space_async(
    space: str,
    *,
    description: str = "",
    parent_space: str = "",
    metadata: dict[str, Any] | None = None,
    update_existing: bool = True,
    conn: asyncpg.Connection | None = None,
) -> str:
    normalized = normalize_space(space)
    if normalized == parent_space:
        parent_space = ""
    own_conn = conn is None
    conn = conn or await connect()
    try:
        if parent_space:
            await ensure_space_async(parent_space, conn=conn)
        if update_existing:
            await conn.execute(
                """
                INSERT INTO agent_memory_spaces (space, description, parent_space, metadata)
                VALUES ($1, $2, $3, $4::JSONB)
                ON CONFLICT (space) DO UPDATE SET
                    description = COALESCE(NULLIF(EXCLUDED.description, ''), agent_memory_spaces.description),
                    parent_space = COALESCE(EXCLUDED.parent_space, agent_memory_spaces.parent_space),
                    metadata = agent_memory_spaces.metadata || EXCLUDED.metadata,
                    updated_at = NOW()
                """,
                normalized,
                description or None,
                normalize_space(parent_space) if parent_space else None,
                metadata or {},
            )
        else:
            await conn.execute(
                """
                INSERT INTO agent_memory_spaces (space, description, parent_space, metadata)
                VALUES ($1, $2, $3, $4::JSONB)
                ON CONFLICT (space) DO NOTHING
                """,
                normalized,
                description or None,
                normalize_space(parent_space) if parent_space else None,
                metadata or {},
            )
        return normalized
    finally:
        if own_conn:
            await conn.close()


async def upsert_entity_async(
    space: str,
    name: str,
    *,
    kind: str = "concept",
    summary: str = "",
    metadata: dict[str, Any] | None = None,
    conn: asyncpg.Connection | None = None,
) -> int | None:
    name = normalize_entity_name(name)
    if not name:
        return None
    own_conn = conn is None
    conn = conn or await connect()
    try:
        row = await conn.fetchrow(
            """
            INSERT INTO agent_memory_entities (space, name, kind, summary, metadata)
            VALUES ($1, $2, $3, $4, $5::JSONB)
            ON CONFLICT (space, lower(name), kind) DO UPDATE SET
                summary = COALESCE(NULLIF(EXCLUDED.summary, ''), agent_memory_entities.summary),
                metadata = agent_memory_entities.metadata || EXCLUDED.metadata,
                updated_at = NOW()
            RETURNING id
            """,
            normalize_space(space),
            name,
            (kind or "concept")[:80],
            summary or None,
            metadata or {},
        )
        return row["id"] if row else None
    finally:
        if own_conn:
            await conn.close()


async def mention_entities_async(
    event_id: int,
    space: str,
    entities: list[dict[str, str]],
    *,
    conn: asyncpg.Connection | None = None,
) -> list[int]:
    own_conn = conn is None
    conn = conn or await connect()
    mentioned: list[int] = []
    try:
        for entity in entities:
            entity_id = await upsert_entity_async(
                space,
                entity.get("name", ""),
                kind=entity.get("kind") or "concept",
                conn=conn,
            )
            if not entity_id:
                continue
            await conn.execute(
                """
                INSERT INTO agent_memory_entity_mentions (event_id, entity_id, weight)
                VALUES ($1, $2, 1.0)
                ON CONFLICT (event_id, entity_id) DO UPDATE SET weight = GREATEST(agent_memory_entity_mentions.weight, 1.0)
                """,
                event_id,
                entity_id,
            )
            mentioned.append(entity_id)
        return mentioned
    finally:
        if own_conn:
            await conn.close()


async def add_event_async(
    *,
    agent: str,
    event_type: str,
    content: str,
    space: str = "",
    memory_kind: str = "event",
    role: str = "",
    source: str = "",
    session_id: str = "",
    summary: str = "",
    tags: list[str] | None = None,
    metadata: dict[str, Any] | None = None,
    event_uid: str = "",
    index_entities: bool = True,
    conn: asyncpg.Connection | None = None,
) -> dict[str, Any]:
    space = normalize_space(space)
    metadata = dict(metadata or {})
    tags = tags or []
    original_content = content
    content, content_archive = compact_indexed_content(original_content)
    if content_archive:
        metadata["content_archive"] = content_archive
    summary = summary or summarize(original_content)
    event_uid = event_uid.strip() or make_uid(agent, event_type, session_id, original_content, metadata, space, memory_kind)
    embedding = vector_literal(cpu_embedding(content + "\n" + summary))

    own_conn = conn is None
    if own_conn:
        conn = await connect()
    assert conn is not None
    transaction = conn.transaction() if own_conn else None
    if transaction is not None:
        await transaction.start()
    try:
        await ensure_space_async(space, update_existing=False, conn=conn)
        row = await conn.fetchrow(
            """
            INSERT INTO agent_memory_events (
                event_uid, space, session_id, agent, event_type, memory_kind, role, source,
                content, summary, tags, metadata, embedding
            )
            VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11::TEXT[], $12::JSONB, $13::VECTOR)
            ON CONFLICT (event_uid) DO UPDATE SET
                metadata = agent_memory_events.metadata || EXCLUDED.metadata,
                space = COALESCE(NULLIF(EXCLUDED.space, 'global'), agent_memory_events.space),
                content = EXCLUDED.content,
                summary = EXCLUDED.summary,
                tags = EXCLUDED.tags,
                embedding = EXCLUDED.embedding
            RETURNING id, created_at, (xmax = 0) AS inserted
            """,
            event_uid,
            space,
            session_id or None,
            agent,
            event_type,
            memory_kind or "event",
            role or None,
            source or None,
            content,
            summary,
            tags,
            metadata,
            embedding,
        )
        mentioned: list[int] = []
        if index_entities:
            entities = extract_entities(content + "\n" + summary, tags, metadata)
            mentioned = await mention_entities_async(row["id"], space, entities, conn=conn)
        archive_result: dict[str, Any] | None = None
        if content_archive:
            archive_result = await archive_payload_async(
                event_id=row["id"],
                event_uid=event_uid,
                payload=original_content,
                payload_kind="oversize_event_content",
                payload_format="text",
                redacted=False,
                space=space,
                session_id=session_id,
                agent=agent,
                event_type=event_type,
                source=source,
                metadata={"payload_sha256": content_archive["payload_sha256"]},
                conn=conn,
            )
        result = {
            "ok": True,
            "id": row["id"],
            "created_at": row["created_at"],
            "inserted": bool(row["inserted"]),
            "event_uid": event_uid,
            "space": space,
            "memory_kind": memory_kind or "event",
            "entities": len(mentioned),
        }
        if archive_result:
            result["content_archive_id"] = archive_result["id"]
        if transaction is not None:
            await transaction.commit()
        return result
    except Exception:
        if transaction is not None:
            await transaction.rollback()
        raise
    finally:
        if own_conn:
            await conn.close()


def archive_text_from_payload(payload: Any) -> str:
    if isinstance(payload, str):
        return payload
    return json.dumps(payload, ensure_ascii=False, default=json_default)


async def archive_payload_async(
    *,
    event_id: int,
    event_uid: str,
    payload: Any,
    payload_kind: str = "hook_payload",
    payload_format: str = "json",
    redacted: bool = True,
    space: str = "",
    session_id: str = "",
    agent: str = "",
    event_type: str = "",
    source: str = "",
    metadata: dict[str, Any] | None = None,
    conn: asyncpg.Connection | None = None,
) -> dict[str, Any]:
    text = archive_text_from_payload(payload)
    raw = text.encode("utf-8", errors="backslashreplace")
    digest = hashlib.sha256(raw).hexdigest()
    compressed = zlib.compress(raw, level=6)
    own_conn = conn is None
    if own_conn:
        conn = await connect()
    assert conn is not None
    try:
        row = await conn.fetchrow(
            """
            INSERT INTO agent_memory_payload_archives (
                event_id, event_uid, payload_sha256, payload_kind, payload_format,
                compression, raw_chars, compressed_bytes, redacted, space, session_id,
                agent, event_type, source, payload, metadata
            )
            VALUES (
                $1, $2, $3, $4, $5,
                'zlib', $6, $7, $8, $9, $10,
                $11, $12, $13, $14::BYTEA, $15::JSONB
            )
            ON CONFLICT (event_uid, payload_kind, payload_sha256) DO UPDATE SET
                event_id = COALESCE(EXCLUDED.event_id, agent_memory_payload_archives.event_id),
                metadata = agent_memory_payload_archives.metadata || EXCLUDED.metadata
            RETURNING id, created_at, (xmax = 0) AS inserted
            """,
            event_id,
            event_uid,
            digest,
            payload_kind or "hook_payload",
            payload_format or "json",
            len(text),
            len(compressed),
            redacted,
            normalize_space(space),
            session_id or None,
            agent or "agent",
            event_type or "event",
            source or None,
            compressed,
            metadata or {},
        )
        return {
            "ok": True,
            "id": row["id"],
            "created_at": row["created_at"],
            "inserted": bool(row["inserted"]),
            "event_id": event_id,
            "event_uid": event_uid,
            "payload_sha256": digest,
            "payload_kind": payload_kind or "hook_payload",
            "raw_chars": len(text),
            "compressed_bytes": len(compressed),
            "redacted": redacted,
        }
    finally:
        if own_conn:
            await conn.close()


async def trace_archives_async(
    *,
    archive_id: int = 0,
    event_id: int = 0,
    event_uid: str = "",
    payload_sha256: str = "",
    session_id: str = "",
    space: str = "",
    all_spaces: bool = False,
    limit: int = 5,
    include_payload: bool = True,
) -> list[dict[str, Any]]:
    filters = []
    params: list[Any] = []
    idx = 1
    if archive_id:
        filters.append(f"id = ${idx}")
        params.append(archive_id)
        idx += 1
    if event_id:
        filters.append(f"event_id = ${idx}")
        params.append(event_id)
        idx += 1
    if event_uid:
        filters.append(f"event_uid = ${idx}")
        params.append(event_uid)
        idx += 1
    if payload_sha256:
        filters.append(f"payload_sha256 = ${idx}")
        params.append(payload_sha256)
        idx += 1
    if session_id:
        filters.append(f"session_id = ${idx}")
        params.append(session_id)
        idx += 1
    if space and not all_spaces:
        filters.append(f"space = ${idx}")
        params.append(normalize_space(space))
        idx += 1
    where = ("WHERE " + " AND ".join(filters)) if filters else ""
    params.append(limit)
    payload_column = ", payload" if include_payload else ""

    conn = await connect()
    try:
        rows = await conn.fetch(
            f"""
            SELECT
                id, event_id, event_uid, payload_sha256, payload_kind, payload_format,
                compression, raw_chars, compressed_bytes, redacted, space, session_id,
                agent, event_type, source, metadata, created_at
                {payload_column}
            FROM agent_memory_payload_archives
            {where}
            ORDER BY created_at DESC
            LIMIT ${idx}
            """,
            *params,
        )
        results: list[dict[str, Any]] = []
        for row in rows:
            item = dict(row)
            if include_payload:
                payload_bytes = bytes(item.pop("payload"))
                if len(payload_bytes) != int(item.get("compressed_bytes") or 0):
                    raise ValueError(f"archive {item['id']} compressed size mismatch")
                if item.get("compression") == "zlib":
                    raw_bytes = zlib.decompress(payload_bytes)
                else:
                    raw_bytes = payload_bytes
                digest = hashlib.sha256(raw_bytes).hexdigest()
                if digest != item.get("payload_sha256"):
                    raise ValueError(f"archive {item['id']} SHA-256 mismatch")
                payload_text = raw_bytes.decode("utf-8", errors="backslashreplace")
                if len(payload_text) != int(item.get("raw_chars") or 0):
                    raise ValueError(f"archive {item['id']} raw character count mismatch")
                item["payload"] = payload_text
                item["integrity_verified"] = True
            results.append(item)
        return results
    finally:
        await conn.close()


async def backfill_legacy_archives_async(*, limit: int = 0, dry_run: bool = False) -> dict[str, Any]:
    limit = max(0, int(limit))
    conn = await connect()
    try:
        candidate_count = await conn.fetchval(
            """
            SELECT count(*)
            FROM agent_memory_events e
            WHERE e.metadata ? 'raw_payload'
              AND NOT EXISTS (
                  SELECT 1 FROM agent_memory_payload_archives a WHERE a.event_id = e.id
              )
            """
        )
        if dry_run:
            return {"ok": True, "dry_run": True, "candidates": candidate_count, "limit": limit}
        limit_sql = "LIMIT $1" if limit else ""
        params: list[Any] = [limit] if limit else []
        inserted = await conn.fetchval(
            f"""
            WITH candidates AS MATERIALIZED (
                SELECT e.*
                FROM agent_memory_events e
                WHERE e.metadata ? 'raw_payload'
                  AND NOT EXISTS (
                      SELECT 1 FROM agent_memory_payload_archives a WHERE a.event_id = e.id
                  )
                ORDER BY e.id
                {limit_sql}
            ), inserted AS (
                INSERT INTO agent_memory_payload_archives (
                    event_id, event_uid, payload_sha256, payload_kind, payload_format,
                    compression, raw_chars, compressed_bytes, redacted, space, session_id,
                    agent, event_type, source, payload, metadata
                )
                SELECT
                    c.id,
                    c.event_uid,
                    encode(sha256(convert_to((c.metadata->'raw_payload')::TEXT, 'UTF8')), 'hex'),
                    'legacy_compact_payload',
                    'json',
                    'none',
                    char_length((c.metadata->'raw_payload')::TEXT),
                    octet_length(convert_to((c.metadata->'raw_payload')::TEXT, 'UTF8')),
                    FALSE,
                    c.space,
                    c.session_id,
                    c.agent,
                    c.event_type,
                    c.source,
                    convert_to((c.metadata->'raw_payload')::TEXT, 'UTF8'),
                    jsonb_build_object(
                        'legacy_source', 'agent_memory_events.metadata.raw_payload',
                        'redaction_version', 'pre-v2-unknown',
                        'full_original_recovered', FALSE
                    )
                FROM candidates c
                ON CONFLICT (event_uid, payload_kind, payload_sha256) DO NOTHING
                RETURNING id
            )
            SELECT count(*) FROM inserted
            """,
            *params,
        )
        return {
            "ok": True,
            "candidates_before": candidate_count,
            "inserted": inserted,
            "limit": limit,
            "remaining_estimate": max(0, int(candidate_count) - int(inserted or 0)),
        }
    finally:
        await conn.close()


async def search_events_async(
    query: str,
    *,
    limit: int = 10,
    space: str = "",
    all_spaces: bool = False,
    agent: str = "",
    event_type: str = "",
    memory_kind: str = "",
    tags: list[str] | None = None,
    min_score: float | None = None,
) -> list[dict[str, Any]]:
    qvec = vector_literal(cpu_embedding(query))
    filters = []
    params: list[Any] = [query, qvec, query]
    idx = 4
    if space and not all_spaces:
        filters.append(f"space = ${idx}")
        params.append(normalize_space(space))
        idx += 1
    if agent:
        filters.append(f"agent = ${idx}")
        params.append(agent)
        idx += 1
    if event_type:
        filters.append(f"event_type = ${idx}")
        params.append(event_type)
        idx += 1
    if memory_kind:
        filters.append(f"memory_kind = ${idx}")
        params.append(memory_kind)
        idx += 1
    if tags:
        filters.append(f"tags && ${idx}::TEXT[]")
        params.append(tags)
        idx += 1
    def where_clause(extra: str = "") -> str:
        parts = [*filters]
        if extra:
            parts.append(extra)
        return ("WHERE " + " AND ".join(parts)) if parts else ""

    where = where_clause()
    text_where = where_clause("search_vector @@ (SELECT tsq FROM q)")
    phrase_where = where_clause("content ILIKE '%' || (SELECT raw_query FROM q) || '%'")
    candidate_limit = search_candidate_limit(limit)
    candidate_idx = idx
    base_params = list(params)
    params.append(candidate_limit)
    idx += 1
    final_limit_idx = idx
    params.append(limit)
    phrase_candidates_enabled = os.getenv("AGENT_MEMORY_ENABLE_PHRASE_CANDIDATES", "").lower() in {"1", "true", "yes"}
    use_phrase_candidates = (
        phrase_candidates_enabled
        and len(query.strip()) >= 12
        and bool(re.search(r"\s", query.strip()))
    )
    phrase_candidate_sql = ""
    phrase_union_sql = ""
    if use_phrase_candidates:
        phrase_candidate_sql = f"""
            ,
            phrase_candidates AS MATERIALIZED (
                SELECT id
                FROM agent_memory_events
                {phrase_where}
                ORDER BY created_at DESC
                LIMIT ${candidate_idx}
            )"""
        phrase_union_sql = """
                UNION
                SELECT id FROM phrase_candidates"""

    threshold = DEFAULT_MIN_SEARCH_SCORE if min_score is None else max(-1.0, min(float(min_score), 1.0))
    cache_key = search_cache_key(
        query,
        limit=limit,
        space=space,
        all_spaces=all_spaces,
        agent=agent,
        event_type=event_type,
        memory_kind=memory_kind,
        tags=tags,
        min_score=threshold,
        candidate_limit=candidate_limit,
        phrase_candidates=use_phrase_candidates,
    )
    identifier_query = bool(re.fullmatch(r"[A-Za-z0-9_.:/-]{8,}", query.strip()))

    def filtered_results(rows: Any) -> list[dict[str, Any]]:
        selected: list[dict[str, Any]] = []
        for row in rows:
            item = dict(row)
            if float(item.get("score") or 0.0) < threshold:
                continue
            if identifier_query and float(item.get("text_score") or 0.0) <= 0.0 and float(item.get("trigram_score") or 0.0) < 0.20:
                continue
            selected.append(item)
        return selected

    def result_sort_key(item: dict[str, Any]) -> tuple[float, float, int]:
        created_at = item.get("created_at")
        if isinstance(created_at, str):
            try:
                created_at = datetime.fromisoformat(created_at.replace("Z", "+00:00"))
            except ValueError:
                created_at = None
        created_timestamp = created_at.timestamp() if isinstance(created_at, datetime) else 0.0
        return float(item.get("score") or 0.0), created_timestamp, int(item.get("id") or 0)

    def merge_results(cached_items: list[dict[str, Any]], new_items: list[dict[str, Any]]) -> list[dict[str, Any]]:
        merged = {int(item["id"]): item for item in cached_items if item.get("id") is not None}
        for item in new_items:
            if item.get("id") is not None:
                merged[int(item["id"])] = item
        return sorted(merged.values(), key=result_sort_key, reverse=True)[:limit]

    conn = await connect()
    try:
        search_timeout_ms = max(1000, min(int(os.getenv("AGENT_MEMORY_SEARCH_TIMEOUT_MS", "15000")), 60000))
        await conn.execute(f"SET statement_timeout = {search_timeout_ms}")
        ef_search = max(40, min(MAX_SEARCH_CANDIDATES, candidate_limit))
        try:
            await conn.execute(f"SET hnsw.ef_search = {ef_search}")
        except Exception:
            pass
        try:
            await conn.execute("SET hnsw.iterative_scan = strict_order")
            await conn.execute(f"SET hnsw.max_scan_tuples = {max(candidate_limit * 20, 20000)}")
        except Exception:
            pass
        async with conn.transaction(isolation="repeatable_read"):
            cache_available = False
            if search_cache_enabled():
                cache_available = bool(
                    await conn.fetchval(
                        """
                        SELECT to_regclass('public.agent_memory_search_cache') IS NOT NULL
                           AND to_regclass('public.agent_memory_search_state') IS NOT NULL
                        """
                    )
                )
            if cache_available:
                state = await conn.fetchrow(
                    """
                    SELECT revision, mutation_revision
                    FROM agent_memory_search_state
                    WHERE singleton = TRUE
                    """
                )
                current_max_event_id = int(
                    await conn.fetchval("SELECT COALESCE(max(id), 0) FROM agent_memory_events") or 0
                )
                cached = await conn.fetchrow(
                    """
                    SELECT revision, mutation_revision, max_event_id, results
                    FROM agent_memory_search_cache cache
                    WHERE cache.cache_key = $1
                    """,
                    cache_key,
                )
                if (
                    cached
                    and state
                    and int(cached["mutation_revision"]) == int(state["mutation_revision"])
                    and int(cached["revision"]) == int(state["revision"])
                    and int(cached["max_event_id"]) == current_max_event_id
                ):
                    await conn.execute(
                        """
                        UPDATE agent_memory_search_cache
                        SET hit_count = hit_count + 1, last_accessed_at = NOW()
                        WHERE cache_key = $1
                        """,
                        cache_key,
                    )
                    return restore_cached_search_results(cached["results"])
                if cached and state and int(cached["mutation_revision"]) == int(state["mutation_revision"]):
                    cached_max_event_id = int(cached["max_event_id"])
                    try:
                        delta_max = max(1, min(int(os.getenv("AGENT_MEMORY_SEARCH_CACHE_DELTA_MAX", "5000")), 50000))
                    except ValueError:
                        delta_max = 5000
                    if (
                        current_max_event_id > cached_max_event_id
                        and current_max_event_id - cached_max_event_id <= delta_max
                    ):
                        delta_start_idx = candidate_idx
                        delta_end_idx = candidate_idx + 1
                        delta_limit_idx = candidate_idx + 2
                        delta_where = where_clause(
                            f"id > ${delta_start_idx} AND id <= ${delta_end_idx}"
                        )
                        delta_rows = await conn.fetch(
                            f"""
                            WITH q AS (
                                SELECT
                                    plainto_tsquery('english', $1) AS tsq,
                                    $2::VECTOR AS emb,
                                    $3 AS raw_query
                            )
                            SELECT
                                e.id, e.created_at, e.space, e.session_id, e.agent, e.event_type, e.memory_kind, e.role, e.source,
                                e.summary, e.tags, e.metadata - 'raw_payload' AS metadata,
                                left(e.content, 2400) AS content,
                                1 - (e.embedding <=> (SELECT emb FROM q)) AS semantic_score,
                                ts_rank_cd(e.search_vector, (SELECT tsq FROM q)) AS text_score,
                                similarity(e.content, (SELECT raw_query FROM q)) AS trigram_score,
                                e.search_vector @@ (SELECT tsq FROM q) AS lexical_match,
                                (
                                    (1 - (e.embedding <=> (SELECT emb FROM q))) * 0.68
                                    + ts_rank_cd(e.search_vector, (SELECT tsq FROM q)) * 0.22
                                    + similarity(e.content, (SELECT raw_query FROM q)) * 0.10
                                    + CASE WHEN e.search_vector @@ (SELECT tsq FROM q) THEN 0.18 ELSE 0.0 END
                                ) AS score
                            FROM agent_memory_events e
                            {delta_where}
                            ORDER BY score DESC, e.created_at DESC, e.id DESC
                            LIMIT ${delta_limit_idx}
                            """,
                            *base_params,
                            cached_max_event_id,
                            current_max_event_id,
                            limit,
                        )
                        merged_results = merge_results(
                            restore_cached_search_results(cached["results"]),
                            filtered_results(delta_rows),
                        )
                        await conn.execute(
                            """
                            UPDATE agent_memory_search_cache
                            SET revision = $2,
                                mutation_revision = $3,
                                max_event_id = $4,
                                results = $5::JSONB,
                                result_count = $6,
                                hit_count = hit_count + 1,
                                refreshed_at = NOW(),
                                last_accessed_at = NOW()
                            WHERE cache_key = $1
                            """,
                            cache_key,
                            int(state["revision"]),
                            int(state["mutation_revision"]),
                            current_max_event_id,
                            merged_results,
                            len(merged_results),
                        )
                        return merged_results

            rows = await conn.fetch(
                f"""
                WITH q AS (
                    SELECT
                        plainto_tsquery('english', $1) AS tsq,
                        $2::VECTOR AS emb,
                        $3 AS raw_query
                ),
                vector_candidates AS MATERIALIZED (
                    SELECT id
                    FROM agent_memory_events
                    {where}
                    ORDER BY embedding <=> (SELECT emb FROM q)
                    LIMIT ${candidate_idx}
                ),
                text_candidates AS MATERIALIZED (
                    SELECT id
                    FROM agent_memory_events
                    {text_where}
                    ORDER BY ts_rank_cd(search_vector, (SELECT tsq FROM q)) DESC, created_at DESC, id DESC
                    LIMIT ${candidate_idx}
                )
                {phrase_candidate_sql},
                candidate_ids AS (
                    SELECT id FROM vector_candidates
                    UNION
                    SELECT id FROM text_candidates
                    {phrase_union_sql}
                )
                SELECT
                    e.id, e.created_at, e.space, e.session_id, e.agent, e.event_type, e.memory_kind, e.role, e.source,
                    e.summary, e.tags, e.metadata - 'raw_payload' AS metadata,
                    left(e.content, 2400) AS content,
                    1 - (e.embedding <=> (SELECT emb FROM q)) AS semantic_score,
                    ts_rank_cd(e.search_vector, (SELECT tsq FROM q)) AS text_score,
                    similarity(e.content, (SELECT raw_query FROM q)) AS trigram_score,
                    e.search_vector @@ (SELECT tsq FROM q) AS lexical_match,
                    (
                        (1 - (e.embedding <=> (SELECT emb FROM q))) * 0.68
                        + ts_rank_cd(e.search_vector, (SELECT tsq FROM q)) * 0.22
                        + similarity(e.content, (SELECT raw_query FROM q)) * 0.10
                        + CASE WHEN e.search_vector @@ (SELECT tsq FROM q) THEN 0.18 ELSE 0.0 END
                    ) AS score
                FROM agent_memory_events e
                JOIN candidate_ids c ON c.id = e.id
                ORDER BY score DESC, e.created_at DESC, e.id DESC
                LIMIT ${final_limit_idx}
                """,
                *params,
            )
            results = filtered_results(rows)

            if cache_available:
                assert state is not None
                await conn.execute(
                    """
                    INSERT INTO agent_memory_search_cache (
                        cache_key, revision, mutation_revision, max_event_id, results, result_count,
                        hit_count, created_at, refreshed_at, last_accessed_at
                    )
                    VALUES ($1, $2, $3, $4, $5::JSONB, $6, 0, NOW(), NOW(), NOW())
                    ON CONFLICT (cache_key) DO UPDATE SET
                        revision = EXCLUDED.revision,
                        mutation_revision = EXCLUDED.mutation_revision,
                        max_event_id = EXCLUDED.max_event_id,
                        results = EXCLUDED.results,
                        result_count = EXCLUDED.result_count,
                        hit_count = 0,
                        refreshed_at = NOW(),
                        last_accessed_at = NOW()
                    """,
                    cache_key,
                    int(state["revision"]),
                    int(state["mutation_revision"]),
                    current_max_event_id,
                    results,
                    len(results),
                )
                await conn.execute(
                    """
                    DELETE FROM agent_memory_search_cache
                    WHERE cache_key IN (
                        SELECT cache_key
                        FROM agent_memory_search_cache
                        WHERE last_accessed_at < NOW() - INTERVAL '7 days'
                        ORDER BY last_accessed_at
                        LIMIT 200
                    )
                    """
                )
            return results
    finally:
        await conn.close()


def parse_datetime_arg(value: str) -> datetime | None:
    text = (value or "").strip()
    if not text:
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        text = text + "T00:00:00+00:00"
    if text.endswith("Z"):
        text = text[:-1] + "+00:00"
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed


async def grep_events_async(
    query: str,
    *,
    limit: int = 20,
    space: str = "",
    all_spaces: bool = False,
    agent: str = "",
    event_type: str = "",
    memory_kind: str = "",
    tags: list[str] | None = None,
    since: str = "",
    contains: bool = False,
    timeout_ms: int = 15000,
) -> list[dict[str, Any]]:
    query = (query or "").strip()
    if not query:
        raise ValueError("query is required")
    limit = max(1, min(int(limit), 500))
    timeout_ms = max(1000, min(int(timeout_ms), 60000))
    filters = []
    params: list[Any] = [query]
    idx = 2
    if space and not all_spaces:
        filters.append(f"space = ${idx}")
        params.append(normalize_space(space))
        idx += 1
    if agent:
        filters.append(f"agent = ${idx}")
        params.append(agent)
        idx += 1
    if event_type:
        filters.append(f"event_type = ${idx}")
        params.append(event_type)
        idx += 1
    if memory_kind:
        filters.append(f"memory_kind = ${idx}")
        params.append(memory_kind)
        idx += 1
    if tags:
        filters.append(f"tags && ${idx}::TEXT[]")
        params.append(tags)
        idx += 1
    since_dt = parse_datetime_arg(since)
    if since_dt:
        filters.append(f"created_at >= ${idx}")
        params.append(since_dt)
        idx += 1
    params.append(limit)
    filter_sql = (" AND " + " AND ".join(filters)) if filters else ""
    conn = await connect()
    try:
        await conn.execute(f"SET statement_timeout = {timeout_ms}")
        if contains:
            contains_candidate_limit = max(200, min(5000, limit * 50))
            contains_params = [*params, contains_candidate_limit]
            contains_candidate_idx = idx + 1
            rows = await conn.fetch(
                f"""
                WITH candidates AS MATERIALIZED (
                    SELECT id
                    FROM agent_memory_events
                    WHERE content ILIKE '%' || $1 || '%'{filter_sql}
                    LIMIT ${contains_candidate_idx}
                )
                SELECT e.id, e.created_at, e.space, e.session_id, e.agent, e.event_type, e.memory_kind, e.role, e.source,
                       e.summary, e.tags, e.metadata - 'raw_payload' AS metadata, left(e.content, 2400) AS content,
                       similarity(e.content, $1) AS score
                FROM agent_memory_events e
                JOIN candidates c ON c.id = e.id
                ORDER BY similarity(e.content, $1) DESC, e.created_at DESC
                LIMIT ${idx}
                """,
                *contains_params,
            )
        else:
            rows = await conn.fetch(
                f"""
                WITH q AS (SELECT websearch_to_tsquery('english', $1) AS tsq)
                SELECT id, created_at, space, session_id, agent, event_type, memory_kind, role, source,
                       summary, tags, metadata - 'raw_payload' AS metadata, left(content, 2400) AS content,
                       ts_rank_cd(search_vector, (SELECT tsq FROM q)) AS score
                FROM agent_memory_events
                WHERE search_vector @@ (SELECT tsq FROM q){filter_sql}
                ORDER BY ts_rank_cd(search_vector, (SELECT tsq FROM q)) DESC, created_at DESC
                LIMIT ${idx}
                """,
                *params,
            )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


def error_signature(error: str) -> str:
    text = re.sub(r"\s+", " ", (error or "").strip())
    text = re.sub(r"\b\d+(?:\.\d+)?\b", "<num>", text)
    text = re.sub(r"0x[0-9a-f]+", "0x<hex>", text, flags=re.IGNORECASE)
    return text[:400] or "unknown"


def summarize_error_log(log_file: str, limit_groups: int = 100) -> dict[str, Any]:
    path = Path(log_file)
    groups: dict[tuple[str, str, str, str, str], dict[str, Any]] = {}
    total = 0
    unreadable = 0
    if not path.exists():
        return {"ok": False, "error": f"log file not found: {path}", "path": str(path), "groups": []}
    with path.open("r", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if not line.strip():
                continue
            total += 1
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                unreadable += 1
                continue
            error = str(record.get("error") or record.get("message") or "unknown")
            error_type = str(record.get("error_type") or "UnknownError")
            event = str(record.get("event") or record.get("hook_event_name") or "agent_memory_cli")
            source = str(record.get("source") or path.stem)
            affected_space = normalize_space(str(record.get("space") or "agent-memory/backend"))
            signature = error_signature(error)
            key = (error_type, signature, event, source, affected_space)
            created_text = str(record.get("created_at") or "")
            created_dt = parse_datetime_arg(created_text)
            stats = groups.setdefault(
                key,
                {
                    "error_type": error_type,
                    "error_signature": signature,
                    "event": event,
                    "source": source,
                    "affected_space": affected_space,
                    "count": 0,
                    "first_seen": created_text,
                    "latest_seen": created_text,
                    "sample_errors": [],
                    "sample_payload_excerpt": "",
                    "payload_sha256": [],
                    "session_ids": [],
                    "duration_ms_max": 0.0,
                    "_first_dt": created_dt,
                    "_latest_dt": created_dt,
                },
            )
            stats["count"] += 1
            current_first = stats.get("_first_dt")
            current_latest = stats.get("_latest_dt")
            if created_dt and (not current_first or created_dt < current_first):
                stats["_first_dt"] = created_dt
                stats["first_seen"] = created_text
            if created_dt and (not current_latest or created_dt > current_latest):
                stats["_latest_dt"] = created_dt
                stats["latest_seen"] = created_text
            if error not in stats["sample_errors"] and len(stats["sample_errors"]) < 3:
                stats["sample_errors"].append(error[:800])
            payload_excerpt = str(record.get("payload_excerpt") or "")
            if payload_excerpt and not stats["sample_payload_excerpt"]:
                stats["sample_payload_excerpt"] = summarize(payload_excerpt, 700)
            payload_hash = str(record.get("payload_sha256") or "")
            if payload_hash and payload_hash not in stats["payload_sha256"] and len(stats["payload_sha256"]) < 8:
                stats["payload_sha256"].append(payload_hash)
            session_id = str(record.get("session_id") or "")
            if session_id and session_id not in stats["session_ids"] and len(stats["session_ids"]) < 8:
                stats["session_ids"].append(session_id)
            try:
                duration = float(record.get("duration_ms") or 0.0)
                stats["duration_ms_max"] = max(float(stats["duration_ms_max"]), duration)
            except (TypeError, ValueError):
                pass
    group_list = list(groups.values())
    for item in group_list:
        item.pop("_first_dt", None)
        item.pop("_latest_dt", None)
    group_list.sort(key=lambda item: (str(item.get("latest_seen") or ""), int(item.get("count") or 0)), reverse=True)
    return {
        "ok": True,
        "path": str(path),
        "total_records": total,
        "unreadable_records": unreadable,
        "group_count": len(group_list),
        "groups": group_list[: max(1, int(limit_groups))],
    }


def format_error_group_content(group: dict[str, Any], source_path: str) -> str:
    lines = [
        "Agent memory hook/CLI error summary for future semantic search and maintenance.",
        f"source_log: {source_path}",
        f"error_type: {group.get('error_type')}",
        f"error_signature: {group.get('error_signature')}",
        f"event: {group.get('event')}",
        f"source: {group.get('source')}",
        f"affected_space: {group.get('affected_space')}",
        f"count: {group.get('count')}",
        f"first_seen: {group.get('first_seen')}",
        f"latest_seen: {group.get('latest_seen')}",
    ]
    samples = group.get("sample_errors") or []
    if samples:
        lines.append("sample_errors:")
        lines.extend(f"- {sample}" for sample in samples)
    if group.get("sample_payload_excerpt"):
        lines.append("sample_payload_excerpt:")
        lines.append(str(group["sample_payload_excerpt"]))
    return "\n".join(lines)


async def ingest_error_log_async(
    *,
    log_file: str = "",
    target_space: str = "agent-memory/backend",
    limit_groups: int = 50,
    dry_run: bool = False,
) -> dict[str, Any]:
    default_log = LOG_DIR / "agent_memory_hook_errors.jsonl"
    path = str(Path(log_file) if log_file else default_log)
    summary = summarize_error_log(path, limit_groups=limit_groups)
    if not summary.get("ok"):
        return summary
    if dry_run:
        summary["dry_run"] = True
        return summary
    results = []
    conn = await connect()
    try:
        async with conn.transaction():
            for group in summary["groups"]:
                content = format_error_group_content(group, path)
                metadata = {
                    "source_log": path,
                    "error_type": group.get("error_type"),
                    "error_signature": group.get("error_signature"),
                    "event": group.get("event"),
                    "source": group.get("source"),
                    "affected_space": group.get("affected_space"),
                    "count": group.get("count"),
                    "first_seen": group.get("first_seen"),
                    "latest_seen": group.get("latest_seen"),
                    "payload_sha256": group.get("payload_sha256"),
                    "session_ids": group.get("session_ids"),
                    "duration_ms_max": group.get("duration_ms_max"),
                }
                tags = [
                    "memory",
                    "memory_error",
                    "hook_error",
                    str(group.get("error_type") or "UnknownError"),
                    str(group.get("event") or "agent_memory_cli"),
                ]
                stable_group_key = json.dumps(
                    {
                        "version": 1,
                        "target_space": normalize_space(target_space),
                        "error_type": group.get("error_type"),
                        "error_signature": group.get("error_signature"),
                        "event": group.get("event"),
                        "source": group.get("source"),
                        "affected_space": group.get("affected_space"),
                    },
                    sort_keys=True,
                    ensure_ascii=True,
                )
                result = await add_event_async(
                    agent="CodexMemoryMaintenance",
                    event_type="hook_error_summary",
                    space=target_space,
                    memory_kind="hook_error_summary",
                    role="system",
                    source="agent_memory_error_log",
                    session_id="agent-memory-error-log",
                    summary=f"{group.get('error_type')} in {group.get('event')} x{group.get('count')}",
                    tags=tags,
                    metadata=metadata,
                    content=content,
                    event_uid=hashlib.sha256(("agent-memory-error-group-v1\x00" + stable_group_key).encode("utf-8")).hexdigest(),
                    index_entities=False,
                    conn=conn,
                )
                results.append(
                    {
                        "id": result.get("id"),
                        "inserted": result.get("inserted"),
                        "error_type": group.get("error_type"),
                        "event": group.get("event"),
                        "count": group.get("count"),
                        "affected_space": group.get("affected_space"),
                    }
                )
    finally:
        await conn.close()
    return {
        "ok": True,
        "path": path,
        "target_space": normalize_space(target_space),
        "total_records": summary.get("total_records"),
        "unreadable_records": summary.get("unreadable_records"),
        "group_count": summary.get("group_count"),
        "ingested_groups": len(results),
        "results": results,
    }


TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json", ".jsonl",
    ".yaml", ".yml", ".xml", ".html", ".htm", ".rtf", ".py", ".ps1", ".sh",
    ".js", ".ts", ".tsx", ".jsx", ".css", ".sql", ".toml", ".ini", ".cfg",
}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", *TEXT_EXTENSIONS}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
AUDIO_EXTENSIONS = {".wav", ".mp3", ".m4a", ".flac", ".ogg", ".aac", ".aiff", ".mid", ".midi"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".webm", ".avi", ".wmv", ".m4v"}
SUPPORTED_ASSET_EXTENSIONS = DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS | AUDIO_EXTENSIONS | VIDEO_EXTENSIONS
SIDECAR_EXTENSIONS = [".txt", ".md", ".markdown", ".srt", ".vtt", ".lrc", ".json", ".csv"]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_text_best_effort(path: Path, max_chars: int = 2_000_000) -> str:
    with path.open("rb") as handle:
        raw = handle.read(max_chars * 4)
    for encoding in ("utf-8", "utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(encoding, errors="strict")[:max_chars]
        except UnicodeError:
            continue
    return raw.decode("utf-8", errors="replace")[:max_chars]


def clean_caption_text(text: str) -> str:
    lines = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.upper() == "WEBVTT":
            continue
        if re.fullmatch(r"\d+", stripped):
            continue
        if "-->" in stripped:
            continue
        lines.append(stripped)
    return "\n".join(lines)


def run_command_text(args: list[str], timeout: int = 30) -> tuple[str, str]:
    try:
        completed = subprocess.run(
            args,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            check=False,
        )
        return completed.stdout or "", completed.stderr or ""
    except Exception as exc:
        return "", str(exc)


def find_executable(name: str) -> str | None:
    found = shutil.which(name)
    if found:
        return found
    local_app_data = os.getenv("LOCALAPPDATA", "").strip()
    if local_app_data:
        candidate = Path(local_app_data) / "Microsoft" / "WinGet" / "Links" / f"{name}.exe"
        if candidate.exists():
            return str(candidate)
    return None


def parse_key_value_lines(text: str) -> dict[str, str]:
    parsed: dict[str, str] = {}
    for line in text.splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = re.sub(r"[^a-z0-9_]+", "_", key.strip().lower()).strip("_")
        if key:
            parsed[key] = value.strip()
    return parsed


def extract_pdf(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    metadata: dict[str, Any] = {"extractor": "pdf"}
    warnings: list[str] = []
    text = ""
    pdfinfo = find_executable("pdfinfo")
    if pdfinfo:
        out, err = run_command_text([pdfinfo, str(path)], timeout=20)
        if out:
            metadata["pdfinfo"] = parse_key_value_lines(out)
        if err.strip():
            warnings.append(f"pdfinfo: {summarize(err, 240)}")
    pdftotext = find_executable("pdftotext")
    if pdftotext:
        with tempfile.TemporaryDirectory(prefix="agent-memory-pdf-") as tmp:
            output = Path(tmp) / "out.txt"
            out, err = run_command_text([pdftotext, "-layout", "-enc", "UTF-8", str(path), str(output)], timeout=60)
            if output.exists():
                text = read_text_best_effort(output)
            if out.strip():
                warnings.append(f"pdftotext stdout: {summarize(out, 240)}")
            if err.strip():
                warnings.append(f"pdftotext: {summarize(err, 240)}")
    if not text.strip():
        try:
            import pypdf  # type: ignore

            reader = pypdf.PdfReader(str(path))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
            metadata["pypdf_pages"] = len(reader.pages)
        except Exception as exc:
            warnings.append(f"pypdf unavailable/failed: {exc}")
    return text, metadata, warnings


def xml_text_from_zip(path: Path, members: list[str], text_tags: set[str]) -> str:
    parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        selected = []
        for member in members:
            if member.endswith("/"):
                selected.extend(name for name in names if name.startswith(member) and name.endswith(".xml"))
            elif member in names:
                selected.append(member)
            else:
                selected.extend(name for name in names if name.startswith(member) and name.endswith(".xml"))
        for name in sorted(set(selected)):
            try:
                root = ET.fromstring(archive.read(name))
            except Exception:
                continue
            current: list[str] = []
            for elem in root.iter():
                tag = elem.tag.rsplit("}", 1)[-1]
                if tag in text_tags and elem.text:
                    current.append(elem.text)
                if tag in {"p", "br", "tr"} and current:
                    joined = " ".join(piece.strip() for piece in current if piece.strip())
                    if joined:
                        parts.append(joined)
                    current = []
            if current:
                joined = " ".join(piece.strip() for piece in current if piece.strip())
                if joined:
                    parts.append(joined)
    return "\n".join(parts)


def extract_docx(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    members = [
        "word/document.xml",
        "word/comments.xml",
        "word/footnotes.xml",
        "word/endnotes.xml",
        "word/header",
        "word/footer",
    ]
    return xml_text_from_zip(path, members, {"t"}), {"extractor": "docx-zip-xml"}, []


def extract_pptx(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    return xml_text_from_zip(path, ["ppt/slides/"], {"t"}), {"extractor": "pptx-zip-xml"}, []


def extract_xlsx(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    warnings: list[str] = []
    rows: list[str] = []
    try:
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value not in (None, "")]
                if values:
                    rows.append("\t".join(values))
        return "\n".join(rows), {"extractor": "openpyxl", "sheets": workbook.sheetnames}, warnings
    except Exception as exc:
        warnings.append(f"openpyxl unavailable/failed: {exc}")
        try:
            text = xml_text_from_zip(path, ["xl/sharedStrings.xml"], {"t"})
            return text, {"extractor": "xlsx-sharedStrings"}, warnings
        except Exception as inner:
            warnings.append(f"xlsx sharedStrings failed: {inner}")
            return "", {"extractor": "xlsx-failed"}, warnings


def extract_csv_like(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    text = read_text_best_effort(path)
    rows: list[str] = []
    try:
        reader = csv.reader(text.splitlines(), delimiter=delimiter)
        for idx, row in enumerate(reader):
            rows.append(" | ".join(cell.strip() for cell in row))
            if idx >= 5000:
                rows.append("<agent-memory csv preview truncated after 5000 rows>")
                break
        return "\n".join(rows), {"extractor": "csv", "delimiter": delimiter}, []
    except Exception as exc:
        return text, {"extractor": "text-fallback"}, [f"csv parse failed: {exc}"]


def extract_image(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    warnings: list[str] = []
    metadata: dict[str, Any] = {"extractor": "pillow-metadata"}
    try:
        from PIL import Image  # type: ignore

        with Image.open(path) as image:
            metadata.update(
                {
                    "format": image.format,
                    "mode": image.mode,
                    "width": image.width,
                    "height": image.height,
                }
            )
            try:
                tiny = image.convert("RGB").resize((1, 1))
                metadata["average_rgb"] = list(tiny.getpixel((0, 0)))
            except Exception as exc:
                warnings.append(f"average color failed: {exc}")
            try:
                exif = image.getexif()
                if exif:
                    metadata["exif_keys"] = [str(key) for key in list(exif.keys())[:40]]
            except Exception as exc:
                warnings.append(f"exif failed: {exc}")
    except Exception as exc:
        warnings.append(f"pillow unavailable/failed: {exc}")
    sidecar_text, sidecars = read_sidecars(path)
    metadata["sidecars"] = sidecars
    description = [
        f"Image asset: {path.name}",
        f"Path: {path}",
    ]
    if sidecar_text.strip():
        description.append("Sidecar text/caption:")
        description.append(sidecar_text)
    description.append(f"Metadata: {json.dumps(metadata, ensure_ascii=False, default=json_default)}")
    return "\n".join(description), metadata, warnings


def extract_ffprobe_metadata(path: Path) -> tuple[dict[str, Any], list[str]]:
    ffprobe = find_executable("ffprobe")
    if not ffprobe:
        return {}, ["ffprobe unavailable"]
    out, err = run_command_text(
        [
            ffprobe,
            "-v",
            "error",
            "-print_format",
            "json",
            "-show_format",
            "-show_streams",
            str(path),
        ],
        timeout=30,
    )
    warnings = [f"ffprobe: {summarize(err, 240)}"] if err.strip() else []
    try:
        return json.loads(out) if out.strip() else {}, warnings
    except json.JSONDecodeError as exc:
        warnings.append(f"ffprobe json parse failed: {exc}")
        return {}, warnings


def searchable_media_metadata(metadata: dict[str, Any]) -> dict[str, Any]:
    summary = {
        key: metadata.get(key)
        for key in ("format", "mode", "width", "height", "channels", "sample_width", "sample_rate", "frames", "duration_seconds")
        if metadata.get(key) is not None
    }
    ffprobe = metadata.get("ffprobe") if isinstance(metadata.get("ffprobe"), dict) else {}
    format_info = ffprobe.get("format") if isinstance(ffprobe.get("format"), dict) else {}
    if format_info:
        summary["container"] = {
            key: format_info.get(key)
            for key in ("format_name", "format_long_name", "duration", "size", "bit_rate")
            if format_info.get(key) is not None
        }
    streams = ffprobe.get("streams") if isinstance(ffprobe.get("streams"), list) else []
    if streams:
        summary["streams"] = [
            {
                key: stream.get(key)
                for key in ("codec_type", "codec_name", "profile", "width", "height", "sample_rate", "channels", "duration", "bit_rate")
                if stream.get(key) is not None
            }
            for stream in streams[:16]
            if isinstance(stream, dict)
        ]
    return summary


def extract_audio(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    metadata: dict[str, Any] = {"extractor": "audio-metadata"}
    warnings: list[str] = []
    if path.suffix.lower() == ".wav":
        try:
            with wave.open(str(path), "rb") as handle:
                frames = handle.getnframes()
                rate = handle.getframerate()
                metadata.update(
                    {
                        "channels": handle.getnchannels(),
                        "sample_width": handle.getsampwidth(),
                        "sample_rate": rate,
                        "frames": frames,
                        "duration_seconds": round(frames / rate, 3) if rate else None,
                    }
                )
        except Exception as exc:
            warnings.append(f"wave metadata failed: {exc}")
    ffmeta, ffwarnings = extract_ffprobe_metadata(path)
    if ffmeta:
        metadata["ffprobe"] = ffmeta
    warnings.extend(ffwarnings)
    sidecar_text, sidecars = read_sidecars(path)
    metadata["sidecars"] = sidecars
    description = [
        f"Audio/music asset: {path.name}",
        f"Path: {path}",
    ]
    if sidecar_text.strip():
        description.append("Sidecar transcript/lyrics/notes:")
        description.append(sidecar_text)
    description.append(f"Metadata: {json.dumps(searchable_media_metadata(metadata), ensure_ascii=False, default=json_default)}")
    return "\n".join(description), metadata, warnings


def extract_video(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    metadata: dict[str, Any] = {"extractor": "video-metadata"}
    ffmeta, warnings = extract_ffprobe_metadata(path)
    if ffmeta:
        metadata["ffprobe"] = ffmeta
    sidecar_text, sidecars = read_sidecars(path)
    metadata["sidecars"] = sidecars
    description = [
        f"Video asset: {path.name}",
        f"Path: {path}",
    ]
    if sidecar_text.strip():
        description.append("Sidecar transcript/captions/notes:")
        description.append(sidecar_text)
    description.append(f"Metadata: {json.dumps(searchable_media_metadata(metadata), ensure_ascii=False, default=json_default)}")
    return "\n".join(description), metadata, warnings


def read_sidecars(path: Path) -> tuple[str, list[dict[str, Any]]]:
    chunks: list[str] = []
    records: list[dict[str, Any]] = []
    for suffix in SIDECAR_EXTENSIONS:
        sidecar = path.with_suffix(suffix)
        if sidecar == path or not sidecar.exists() or not sidecar.is_file():
            continue
        text = read_text_best_effort(sidecar, max_chars=500_000)
        if suffix in {".srt", ".vtt", ".lrc"}:
            text = clean_caption_text(text)
        chunks.append(f"[sidecar {sidecar.name}]\n{text}")
        records.append(
            {
                "path": str(sidecar),
                "suffix": suffix,
                "size": sidecar.stat().st_size,
                "sha256": file_sha256(sidecar),
            }
        )
    return "\n\n".join(chunks), records


def asset_modality(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in IMAGE_EXTENSIONS:
        return "image"
    if suffix in VIDEO_EXTENSIONS:
        return "video"
    if suffix in AUDIO_EXTENSIONS:
        return "audio"
    if suffix in DOCUMENT_EXTENSIONS:
        return "document"
    return "file"


def extract_asset_text(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    warnings: list[str] = []
    metadata: dict[str, Any] = {}
    if suffix == ".pdf":
        text, metadata, warnings = extract_pdf(path)
    elif suffix == ".docx":
        text, metadata, warnings = extract_docx(path)
    elif suffix == ".pptx":
        text, metadata, warnings = extract_pptx(path)
    elif suffix == ".xlsx":
        text, metadata, warnings = extract_xlsx(path)
    elif suffix in {".csv", ".tsv"}:
        text, metadata, warnings = extract_csv_like(path)
    elif suffix in TEXT_EXTENSIONS:
        text, metadata, warnings = read_text_best_effort(path), {"extractor": "plain-text"}, []
    elif suffix in IMAGE_EXTENSIONS:
        text, metadata, warnings = extract_image(path)
    elif suffix in AUDIO_EXTENSIONS:
        text, metadata, warnings = extract_audio(path)
    elif suffix in VIDEO_EXTENSIONS:
        text, metadata, warnings = extract_video(path)
    else:
        sidecar_text, sidecars = read_sidecars(path)
        metadata = {"extractor": "metadata-only", "sidecars": sidecars}
        text = sidecar_text or f"File asset: {path.name}\nPath: {path}"
    text = text.replace("\x00", "\\x00")
    return {
        "text": text,
        "metadata": metadata,
        "warnings": warnings,
        "modality": asset_modality(path),
        "mime_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
    }


def chunk_text(text: str, *, chunk_chars: int = 3500, overlap: int = 350) -> list[dict[str, Any]]:
    clean = re.sub(r"\r\n?", "\n", text or "").strip()
    if not clean:
        return []
    chunk_chars = max(500, min(int(chunk_chars), 20000))
    overlap = max(0, min(int(overlap), chunk_chars // 2))
    chunks = []
    start = 0
    while start < len(clean):
        end = min(len(clean), start + chunk_chars)
        if end < len(clean):
            boundary = max(clean.rfind("\n\n", start, end), clean.rfind(". ", start, end))
            if boundary > start + chunk_chars // 2:
                end = boundary + 1
        chunks.append({"index": len(chunks), "start": start, "end": end, "text": clean[start:end].strip()})
        if end >= len(clean):
            break
        start = max(end - overlap, start + 1)
    return chunks


def parse_extension_filter(value: str) -> set[str]:
    if not value:
        return set(SUPPORTED_ASSET_EXTENSIONS)
    result = set()
    for item in value.split(","):
        item = item.strip().lower()
        if not item:
            continue
        if not item.startswith("."):
            item = "." + item
        result.add(item)
    return result


def asset_summary_text(path: Path, file_meta: dict[str, Any], extracted: dict[str, Any], text: str, chunk_count: int) -> str:
    preview = summarize(text, 1200)
    return "\n".join(
        [
            "Document/media memory asset.",
            f"name: {path.name}",
            f"path: {path}",
            f"modality: {extracted.get('modality')}",
            f"mime_type: {extracted.get('mime_type')}",
            f"file_sha256: {file_meta.get('file_sha256')}",
            f"chunks: {chunk_count}",
            f"extractor_metadata: {json.dumps(extracted.get('metadata') or {}, ensure_ascii=False, default=json_default)}",
            "preview:",
            preview,
        ]
    )


async def ingest_file_async(
    file_path: str,
    *,
    space: str = "documents/personal",
    agent: str = "Codex",
    source: str = "agent_memory_ingest",
    session_id: str = "",
    tags: list[str] | None = None,
    chunk_chars: int = 3500,
    chunk_overlap: int = 350,
    archive: bool = True,
    dry_run: bool = False,
) -> dict[str, Any]:
    path = Path(file_path).expanduser().resolve()
    if not path.exists() or not path.is_file():
        return {"ok": False, "error": f"file not found: {path}"}
    stat_before = path.stat()
    extracted = extract_asset_text(path)
    digest = file_sha256(path)
    stat = path.stat()
    if stat.st_size != stat_before.st_size or stat.st_mtime_ns != stat_before.st_mtime_ns:
        return {"ok": False, "error": f"file changed during extraction: {path}"}
    file_meta = {
        "path": str(path),
        "name": path.name,
        "suffix": path.suffix.lower(),
        "file_size": stat.st_size,
        "file_mtime": datetime.fromtimestamp(stat.st_mtime, UTC).isoformat(),
        "file_sha256": digest,
    }
    text = extracted.get("text") or ""
    if not text.strip():
        text = f"Document/media asset with no extractable text: {path.name}\nPath: {path}"
    chunks = chunk_text(text, chunk_chars=chunk_chars, overlap=chunk_overlap)
    metadata = {
        **file_meta,
        "modality": extracted.get("modality"),
        "mime_type": extracted.get("mime_type"),
        "extractor": (extracted.get("metadata") or {}).get("extractor"),
        "extractor_metadata": extracted.get("metadata") or {},
        "warnings": extracted.get("warnings") or [],
        "chunk_chars": chunk_chars,
        "chunk_overlap": chunk_overlap,
        "chunk_count": len(chunks),
    }
    base_tags = list(dict.fromkeys([*(tags or []), "document_memory", str(extracted.get("modality")), path.suffix.lower().lstrip(".")]))
    if dry_run:
        return {
            "ok": True,
            "dry_run": True,
            "path": str(path),
            "metadata": metadata,
            "chunk_count": len(chunks),
            "preview": summarize(text, 800),
        }
    conn = await connect()
    try:
        async with conn.transaction():
            asset = await add_event_async(
                agent=agent,
                event_type="document_ingest",
                space=space,
                memory_kind="document_asset",
                role="system",
                source=source,
                session_id=session_id,
                summary=f"{extracted.get('modality')} asset {path.name} ({len(chunks)} chunks)",
                tags=base_tags,
                metadata=metadata,
                content=asset_summary_text(path, file_meta, extracted, text, len(chunks)),
                conn=conn,
            )
            archive_result: dict[str, Any] | None = None
            if archive:
                archive_result = await archive_payload_async(
                    event_id=asset["id"],
                    event_uid=asset["event_uid"],
                    payload={"text": text, "metadata": metadata},
                    payload_kind="document_extraction",
                    payload_format="json",
                    redacted=False,
                    space=space,
                    session_id=session_id,
                    agent=agent,
                    event_type="document_ingest",
                    source=source,
                    metadata={"file_sha256": file_meta["file_sha256"], "path": str(path)},
                    conn=conn,
                )
            chunk_results = []
            total = len(chunks)
            for chunk in chunks:
                chunk_metadata = {
                    **file_meta,
                    "asset_event_id": asset["id"],
                    "modality": extracted.get("modality"),
                    "mime_type": extracted.get("mime_type"),
                    "chunk_index": chunk["index"],
                    "chunk_total": total,
                    "char_start": chunk["start"],
                    "char_end": chunk["end"],
                }
                chunk_content = "\n".join(
                    [
                        f"Document/media chunk {chunk['index'] + 1}/{total}",
                        f"name: {path.name}",
                        f"path: {path}",
                        f"modality: {extracted.get('modality')}",
                        "",
                        chunk["text"],
                    ]
                )
                row = await add_event_async(
                    agent=agent,
                    event_type="document_chunk",
                    space=space,
                    memory_kind="document_chunk",
                    role="system",
                    source=source,
                    session_id=session_id,
                    summary=f"{path.name} chunk {chunk['index'] + 1}/{total}",
                    tags=base_tags,
                    metadata=chunk_metadata,
                    content=chunk_content,
                    conn=conn,
                )
                chunk_results.append({"id": row["id"], "inserted": row["inserted"], "chunk_index": chunk["index"]})
        return {
            "ok": True,
            "path": str(path),
            "space": normalize_space(space),
            "asset_event_id": asset["id"],
            "asset_inserted": asset["inserted"],
            "archive_id": archive_result.get("id") if archive_result else None,
            "archive_inserted": archive_result.get("inserted") if archive_result else None,
            "chunk_count": total,
            "chunks": chunk_results,
            "metadata": metadata,
        }
    finally:
        await conn.close()


async def ingest_path_async(
    root_path: str,
    *,
    recursive: bool = False,
    extensions: str = "",
    max_files: int = 100,
    **kwargs: Any,
) -> dict[str, Any]:
    root = Path(root_path).expanduser().resolve()
    allowed = parse_extension_filter(extensions)
    if root.is_file():
        return await ingest_file_async(str(root), **kwargs)
    if not root.exists() or not root.is_dir():
        return {"ok": False, "error": f"path not found: {root}"}
    iterator = root.rglob("*") if recursive else root.glob("*")
    files = [path for path in iterator if path.is_file() and path.suffix.lower() in allowed]
    files = sorted(files)[: max(1, int(max_files))]
    results = []
    for path in files:
        results.append(await ingest_file_async(str(path), **kwargs))
    return {
        "ok": True,
        "root": str(root),
        "recursive": recursive,
        "matched_files": len(files),
        "results": results,
    }


async def audit_events_async(
    limit: int = 20,
    session_id: str = "",
    agent: str = "",
    space: str = "",
    all_spaces: bool = False,
    model_id: str = "",
    harness: str = "",
    thread_id: str = "",
    provider: str = "",
) -> list[dict[str, Any]]:
    filters = []
    params: list[Any] = []
    idx = 1
    if session_id:
        filters.append(f"session_id = ${idx}")
        params.append(session_id)
        idx += 1
    if agent:
        filters.append(f"agent = ${idx}")
        params.append(agent)
        idx += 1
    if space and not all_spaces:
        filters.append(f"space = ${idx}")
        params.append(normalize_space(space))
        idx += 1
    identity_filter = {key: value for key, value in (("model_id", model_id), ("harness", harness), ("thread_id", thread_id), ("model_provider", provider)) if value}
    if identity_filter:
        filters.append(f"metadata ? 'identity' AND metadata->'identity' @> ${idx}::jsonb")
        params.append(identity_filter)
        idx += 1
    where = ("WHERE " + " AND ".join(filters)) if filters else ""
    params.append(limit)

    conn = await connect()
    try:
        rows = await conn.fetch(
            f"""
            SELECT id, created_at, space, session_id, agent, event_type, memory_kind, role, source,
                   summary, tags, metadata - 'raw_payload' AS metadata, left(content, 2400) AS content
            FROM agent_memory_events
            {where}
            ORDER BY created_at DESC
            LIMIT ${idx}
            """,
            *params,
        )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def status_async(*, details: bool = False) -> dict[str, Any]:
    conn = await connect()
    try:
        total = await conn.fetchval("SELECT count(*) AS total FROM agent_memory_events")
        space_count = await conn.fetchval("SELECT count(*) FROM agent_memory_spaces")
        top_space_rows = await conn.fetch(
            """
            SELECT space, count(*) AS events
            FROM agent_memory_events
            GROUP BY space
            ORDER BY count(*) DESC, space
            LIMIT 20
            """
        )
        graph_counts = await conn.fetchrow(
            """
            SELECT
                (SELECT count(*) FROM agent_memory_entities) AS entities,
                (SELECT count(*) FROM agent_memory_entity_edges) AS entity_edges,
                (SELECT count(*) FROM agent_memory_event_links) AS event_links,
                (SELECT count(*) FROM agent_memory_payload_archives) AS payload_archives,
                (SELECT COALESCE(sum(raw_chars), 0) FROM agent_memory_payload_archives) AS payload_archive_raw_chars,
                (SELECT COALESCE(sum(compressed_bytes), 0) FROM agent_memory_payload_archives) AS payload_archive_compressed_bytes
            """
        )
        span = await conn.fetchrow(
            """
            SELECT max(created_at) AS newest, min(created_at) AS oldest
            FROM agent_memory_events
            """
        )
        result = {
            "ok": True,
            "driver": "asyncpg",
            "total": total,
            "span": dict(span) if span else {},
            "space_count": space_count,
            "top_spaces": [dict(row) for row in top_space_rows],
            "graph": dict(graph_counts) if graph_counts else {},
        }
        cache_available = bool(
            await conn.fetchval(
                """
                SELECT to_regclass('public.agent_memory_search_cache') IS NOT NULL
                   AND to_regclass('public.agent_memory_search_state') IS NOT NULL
                """
            )
        )
        if cache_available:
            cache_stats = await conn.fetchrow(
                """
                SELECT
                    state.revision,
                    state.mutation_revision,
                    count(cache.cache_key) AS entries,
                    COALESCE(sum(cache.hit_count), 0) AS hits,
                    max(cache.last_accessed_at) AS last_accessed_at
                FROM agent_memory_search_state state
                LEFT JOIN agent_memory_search_cache cache ON TRUE
                WHERE state.singleton = TRUE
                GROUP BY state.revision, state.mutation_revision
                """
            )
            result["search_cache"] = dict(cache_stats) if cache_stats else {"revision": 0, "mutation_revision": 0, "entries": 0, "hits": 0}
        else:
            result["search_cache"] = {"available": False}
        if details:
            breakdown_rows = await conn.fetch(
                """
                SELECT space, agent, event_type, count(*) AS count
                FROM agent_memory_events
                GROUP BY space, agent, event_type
                ORDER BY space, agent, event_type
                """
            )
            space_rows = await conn.fetch(
                """
                SELECT s.space, s.description, s.parent_space, count(e.id) AS events
                FROM agent_memory_spaces s
                LEFT JOIN agent_memory_events e ON e.space = s.space
                GROUP BY s.space, s.description, s.parent_space
                ORDER BY s.space
                """
            )
            result["spaces"] = [dict(row) for row in space_rows]
            result["breakdown"] = [dict(row) for row in breakdown_rows]
        return result
    finally:
        await conn.close()


async def list_spaces_async() -> list[dict[str, Any]]:
    conn = await connect()
    try:
        rows = await conn.fetch(
            """
            SELECT s.*, count(e.id) AS event_count
            FROM agent_memory_spaces s
            LEFT JOIN agent_memory_events e ON e.space = s.space
            GROUP BY s.space, s.description, s.parent_space, s.metadata, s.created_at, s.updated_at
            ORDER BY s.updated_at DESC, s.space
            """
        )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def search_entities_async(query: str = "", *, space: str = "", all_spaces: bool = False, limit: int = 20) -> list[dict[str, Any]]:
    filters = []
    params: list[Any] = []
    idx = 1
    if query:
        filters.append(f"(name ILIKE ${idx} OR summary ILIKE ${idx})")
        params.append(f"%{query}%")
        idx += 1
    if space and not all_spaces:
        filters.append(f"space = ${idx}")
        params.append(normalize_space(space))
        idx += 1
    where = ("WHERE " + " AND ".join(filters)) if filters else ""
    params.append(limit)
    conn = await connect()
    try:
        rows = await conn.fetch(
            f"""
            SELECT e.*, count(m.event_id) AS mentions
            FROM agent_memory_entities e
            LEFT JOIN agent_memory_entity_mentions m ON m.entity_id = e.id
            {where}
            GROUP BY e.id
            ORDER BY mentions DESC, e.updated_at DESC
            LIMIT ${idx}
            """,
            *params,
        )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


async def relate_entities_async(
    *,
    source: str,
    target: str,
    relation: str = "related_to",
    space: str = "",
    description: str = "",
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    space = normalize_space(space)
    conn = await connect()
    try:
        await ensure_space_async(space, conn=conn)
        source_id = await upsert_entity_async(space, source, conn=conn)
        target_id = await upsert_entity_async(space, target, conn=conn)
        if not source_id or not target_id:
            return {"ok": False, "error": "source and target are required"}
        row = await conn.fetchrow(
            """
            INSERT INTO agent_memory_entity_edges (
                space, source_entity_id, target_entity_id, relation, description, metadata
            )
            VALUES ($1, $2, $3, $4, $5, $6::JSONB)
            ON CONFLICT (space, source_entity_id, target_entity_id, relation) DO UPDATE SET
                description = COALESCE(NULLIF(EXCLUDED.description, ''), agent_memory_entity_edges.description),
                metadata = agent_memory_entity_edges.metadata || EXCLUDED.metadata,
                updated_at = NOW()
            RETURNING id
            """,
            space,
            source_id,
            target_id,
            relation or "related_to",
            description or None,
            metadata or {},
        )
        return {"ok": True, "id": row["id"], "space": space, "source_entity_id": source_id, "target_entity_id": target_id}
    finally:
        await conn.close()


async def graph_neighbors_async(entity: str, *, space: str = "", all_spaces: bool = False, limit: int = 30) -> list[dict[str, Any]]:
    filters = ["lower(src.name) = lower($1)"]
    params: list[Any] = [entity]
    idx = 2
    if space and not all_spaces:
        filters.append(f"edge.space = ${idx}")
        params.append(normalize_space(space))
        idx += 1
    where = "WHERE " + " AND ".join(filters)
    params.append(limit)
    conn = await connect()
    try:
        rows = await conn.fetch(
            f"""
            SELECT edge.id, edge.space, src.name AS source, edge.relation,
                   dst.name AS target, edge.description, edge.metadata, edge.updated_at
            FROM agent_memory_entity_edges edge
            JOIN agent_memory_entities src ON src.id = edge.source_entity_id
            JOIN agent_memory_entities dst ON dst.id = edge.target_entity_id
            {where}
            ORDER BY edge.updated_at DESC
            LIMIT ${idx}
            """,
            *params,
        )
        return [dict(row) for row in rows]
    finally:
        await conn.close()


def init_db() -> dict[str, Any]:
    return asyncio.run(init_db_async())


def add_event(**kwargs: Any) -> dict[str, Any]:
    return asyncio.run(add_event_async(**kwargs))


def search_events(query: str, **kwargs: Any) -> list[dict[str, Any]]:
    return asyncio.run(search_events_async(query, **kwargs))


def grep_events(query: str, **kwargs: Any) -> list[dict[str, Any]]:
    return asyncio.run(grep_events_async(query, **kwargs))


def audit_events(limit: int = 20, session_id: str = "", agent: str = "", space: str = "", all_spaces: bool = False) -> list[dict[str, Any]]:
    return asyncio.run(audit_events_async(limit, session_id, agent, space, all_spaces))


def status() -> dict[str, Any]:
    return asyncio.run(status_async())


def list_spaces() -> list[dict[str, Any]]:
    return asyncio.run(list_spaces_async())


def search_entities(query: str = "", **kwargs: Any) -> list[dict[str, Any]]:
    return asyncio.run(search_entities_async(query, **kwargs))


def relate_entities(**kwargs: Any) -> dict[str, Any]:
    return asyncio.run(relate_entities_async(**kwargs))


def graph_neighbors(entity: str, **kwargs: Any) -> list[dict[str, Any]]:
    return asyncio.run(graph_neighbors_async(entity, **kwargs))


def read_content(args: argparse.Namespace) -> str:
    if getattr(args, "content_file", ""):
        return Path(args.content_file).read_text(encoding="utf-8")
    if getattr(args, "stdin", False):
        return sys.stdin.read()
    return getattr(args, "content", "") or ""


def parse_json_arg(value: str, default: Any) -> Any:
    if not value:
        return default
    try:
        return json.loads(value)
    except json.JSONDecodeError:
        try:
            return json.loads(value.replace('\\"', '"'))
        except json.JSONDecodeError:
            pass
    stripped = value.strip()
    if stripped.startswith("{") and stripped.endswith("}"):
        pairs = stripped[1:-1].split(",")
        parsed: dict[str, Any] = {}
        for pair in pairs:
            if ":" not in pair:
                return default
            key, raw_val = pair.split(":", 1)
            key = key.strip().strip("'\"")
            raw_val = raw_val.strip().strip("'\"")
            if not key:
                return default
            lowered = raw_val.lower()
            if lowered == "true":
                parsed[key] = True
            elif lowered == "false":
                parsed[key] = False
            elif lowered == "null":
                parsed[key] = None
            else:
                try:
                    parsed[key] = int(raw_val)
                except ValueError:
                    try:
                        parsed[key] = float(raw_val)
                    except ValueError:
                        parsed[key] = raw_val
        return parsed
    return default


def print_result(data: Any, as_json: bool = False) -> None:
    if as_json:
        print(json.dumps(data, indent=2, ensure_ascii=False, default=json_default))
        return
    if isinstance(data, list):
        for item in data:
            created = item.get("created_at")
            score = item.get("score")
            score_text = f" score={score:.3f}" if isinstance(score, float) else ""
            space_text = f" space={item.get('space')}" if item.get("space") else ""
            identity = (item.get("metadata") or {}).get("identity", {})
            identity_text = "".join(f" {key}={identity[key]}" for key in ("model_id", "harness", "thread_id", "agent_path", "model_status") if identity.get(key))
            print(f"[{item.get('id')}] {created}{space_text} {item.get('agent')}/{item.get('event_type')}{score_text}{identity_text}")
            print(f"  {item.get('summary')}")
            print()
        return
    print(json.dumps(data, indent=2, ensure_ascii=False, default=json_default))


async def self_test_async() -> dict[str, Any]:
    await init_db_async()
    sentinel = "AGENT_MEMORY_SELF_TEST_20260512 asyncpg cpu embedder postgres semantic audit trace"
    space_a = "self-test-alpha"
    space_b = "self-test-beta"
    added = await add_event_async(
        agent="self_test",
        event_type="diagnostic",
        space=space_a,
        memory_kind="test",
        role="system",
        source="agent_memory.py",
        session_id="self-test",
        content=sentinel,
        tags=["self_test", "memory", "asyncpg"],
        metadata={"test": True, "driver": "asyncpg", "entities": [{"name": "Agent Memory Spaces", "kind": "concept"}]},
    )
    await add_event_async(
        agent="self_test",
        event_type="diagnostic",
        space=space_b,
        memory_kind="test",
        role="system",
        source="agent_memory.py",
        session_id="self-test",
        content="AGENT_MEMORY_SELF_TEST_20260512 beta isolated unrelated context marker",
        tags=["self_test", "beta"],
        metadata={"test": True, "entities": ["Memory Space Isolation"]},
    )
    relation = await relate_entities_async(
        source="Agent Memory Spaces",
        target="Memory Space Isolation",
        relation="mitigates",
        space=space_a,
        description="Spaces prevent unrelated ideas and workflows from merging in default retrieval.",
        metadata={"test": True},
    )
    results = await search_events_async("asyncpg semantic audit trace memory self test", space=space_a, limit=5)
    beta_results = await search_events_async("beta isolated unrelated context marker", space=space_a, limit=5)
    graph = await graph_neighbors_async("Agent Memory Spaces", space=space_a, limit=5)
    found = any(sentinel in r.get("content", "") for r in results)
    isolated = not any("beta isolated" in r.get("content", "") for r in beta_results)
    return {"ok": found and isolated and bool(graph), "added": added, "relation": relation, "isolated": isolated, "top_results": results, "graph": graph}


def self_test() -> dict[str, Any]:
    return asyncio.run(self_test_async())


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="PostgreSQL-only async agent memory")
    parser.add_argument("--json", action="store_true", help="Emit JSON")
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("init", help="Initialize schema")
    status_parser = sub.add_parser("status", help="Show bounded memory status")
    status_parser.add_argument("--details", action="store_true", help="Include every space and agent/event breakdown row")
    sub.add_parser("self-test", help="Run an insert/search test")

    add = sub.add_parser("add", help="Add a memory event")
    add.add_argument("--agent", default=os.getenv("AGENT_MEMORY_AGENT", "agent"))
    add.add_argument("--event-type", default="observation")
    add.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", DEFAULT_SPACE))
    add.add_argument("--memory-kind", default="event")
    add.add_argument("--role", default="")
    add.add_argument("--source", default="")
    add.add_argument("--session-id", default=os.getenv("AGENT_MEMORY_SESSION_ID", ""))
    add.add_argument("--summary", default="")
    add.add_argument("--tags", default="")
    add.add_argument("--metadata", default="")
    add.add_argument("--model-id", default="")
    add.add_argument("--provider", default="")
    add.add_argument("--harness", default="")
    add.add_argument("--thread-id", default="")
    add.add_argument("--agent-id", default="")
    add.add_argument("--reasoning-effort", default="")
    add.add_argument("--content", default="")
    add.add_argument("--content-file", default="")
    add.add_argument("--stdin", action="store_true")

    search = sub.add_parser("search", help="Search memory")
    search.add_argument("query")
    search.add_argument("--limit", type=int, default=10)
    search.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", ""))
    search.add_argument("--all-spaces", action="store_true")
    search.add_argument("--agent", default="")
    search.add_argument("--event-type", default="")
    search.add_argument("--memory-kind", default="")
    search.add_argument("--tags", default="")
    search.add_argument("--min-score", type=float, default=None, help="Minimum blended retrieval score; defaults to AGENT_MEMORY_MIN_SEARCH_SCORE")

    grep = sub.add_parser("grep", help="Indexed exact/full-text memory search for audits")
    grep.add_argument("query")
    grep.add_argument("--limit", type=int, default=20)
    grep.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", ""))
    grep.add_argument("--all-spaces", action="store_true")
    grep.add_argument("--agent", default="")
    grep.add_argument("--event-type", default="")
    grep.add_argument("--memory-kind", default="")
    grep.add_argument("--tags", default="")
    grep.add_argument("--since", default="")
    grep.add_argument("--contains", action="store_true", help="Use indexed trigram substring matching instead of full-text matching")
    grep.add_argument("--timeout-ms", type=int, default=15000)

    audit = sub.add_parser("audit", help="Show recent audit records")
    audit.add_argument("--limit", type=int, default=20)
    audit.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", ""))
    audit.add_argument("--all-spaces", action="store_true")
    audit.add_argument("--session-id", default="")
    audit.add_argument("--agent", default="")
    audit.add_argument("--model-id", default="", help="Exact model ID in identity metadata")
    audit.add_argument("--harness", default="")
    audit.add_argument("--thread-id", default="")
    audit.add_argument("--provider", default="")

    trace = sub.add_parser("trace", help="Retrieve compressed full payload archives")
    trace.add_argument("--archive-id", type=int, default=0)
    trace.add_argument("--event-id", type=int, default=0)
    trace.add_argument("--event-uid", default="")
    trace.add_argument("--sha256", default="")
    trace.add_argument("--session-id", default="")
    trace.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", ""))
    trace.add_argument("--all-spaces", action="store_true")
    trace.add_argument("--limit", type=int, default=5)
    trace.add_argument("--no-payload", action="store_true", help="Return archive metadata without decompressed payload text")
    trace.add_argument("--content-only", action="store_true", help="Print only decompressed payload text")

    legacy_archives = sub.add_parser("backfill-legacy-archives", help="Archive legacy metadata.raw_payload rows missing trace archives")
    legacy_archives.add_argument("--limit", type=int, default=0, help="Maximum rows to backfill; zero means all")
    legacy_archives.add_argument("--dry-run", action="store_true")

    identities = sub.add_parser("backfill-identities", help="Preview or apply bounded transcript-based model/harness identity enrichment")
    identities.add_argument("--limit", type=int, default=1000)
    identities.add_argument("--before-id", type=int, default=0)
    identities.add_argument("--since", default="")
    identities.add_argument("--apply", action="store_true", help="Write only missing identity metadata; retain event content and times")

    codex_import = sub.add_parser("ingest-codex-transcript", help="Recover missed audit hooks from one exact native Codex transcript")
    codex_import.add_argument("path")
    codex_import.add_argument("--limit", type=int, default=500)
    codex_import.add_argument("--apply", action="store_true")

    errors = sub.add_parser("ingest-error-log", help="Ingest local memory hook/CLI error logs as searchable summaries")
    errors.add_argument("--log-file", default=str(LOG_DIR / "agent_memory_hook_errors.jsonl"))
    errors.add_argument("--target-space", default="agent-memory/backend")
    errors.add_argument("--limit-groups", type=int, default=50)
    errors.add_argument("--dry-run", action="store_true")

    ingest_file = sub.add_parser("ingest-file", help="Extract and store one document/media file as searchable memory chunks")
    ingest_file.add_argument("path")
    ingest_file.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", "documents/personal"))
    ingest_file.add_argument("--agent", default=os.getenv("AGENT_MEMORY_AGENT", "Codex"))
    ingest_file.add_argument("--source", default="agent_memory_ingest")
    ingest_file.add_argument("--session-id", default=os.getenv("AGENT_MEMORY_SESSION_ID", ""))
    ingest_file.add_argument("--tags", default="")
    ingest_file.add_argument("--chunk-chars", type=int, default=3500)
    ingest_file.add_argument("--chunk-overlap", type=int, default=350)
    ingest_file.add_argument("--no-archive", action="store_true")
    ingest_file.add_argument("--dry-run", action="store_true")

    ingest_path = sub.add_parser("ingest-path", help="Extract and store supported files from a directory")
    ingest_path.add_argument("path")
    ingest_path.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", "documents/personal"))
    ingest_path.add_argument("--agent", default=os.getenv("AGENT_MEMORY_AGENT", "Codex"))
    ingest_path.add_argument("--source", default="agent_memory_ingest")
    ingest_path.add_argument("--session-id", default=os.getenv("AGENT_MEMORY_SESSION_ID", ""))
    ingest_path.add_argument("--tags", default="")
    ingest_path.add_argument("--chunk-chars", type=int, default=3500)
    ingest_path.add_argument("--chunk-overlap", type=int, default=350)
    ingest_path.add_argument("--no-archive", action="store_true")
    ingest_path.add_argument("--dry-run", action="store_true")
    ingest_path.add_argument("--recursive", action="store_true")
    ingest_path.add_argument("--extensions", default="")
    ingest_path.add_argument("--max-files", type=int, default=100)

    spaces = sub.add_parser("spaces", help="List or upsert memory spaces")
    spaces.add_argument("--space", default="")
    spaces.add_argument("--description", default="")
    spaces.add_argument("--parent-space", default="")
    spaces.add_argument("--metadata", default="")

    entities = sub.add_parser("entities", help="List/search graph entities")
    entities.add_argument("--query", default="")
    entities.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", ""))
    entities.add_argument("--all-spaces", action="store_true")
    entities.add_argument("--limit", type=int, default=20)

    relate = sub.add_parser("relate", help="Create an entity relationship")
    relate.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", DEFAULT_SPACE))
    relate.add_argument("--source", required=True)
    relate.add_argument("--target", required=True)
    relate.add_argument("--relation", default="related_to")
    relate.add_argument("--description", default="")
    relate.add_argument("--metadata", default="")

    graph = sub.add_parser("graph", help="Show graph neighbors for an entity")
    graph.add_argument("entity")
    graph.add_argument("--space", default=os.getenv("AGENT_MEMORY_SPACE", ""))
    graph.add_argument("--all-spaces", action="store_true")
    graph.add_argument("--limit", type=int, default=30)

    return parser


async def run_cli(args: argparse.Namespace) -> int:
    if args.command == "init":
        print_result(await init_db_async(), args.json)
    elif args.command == "status":
        print_result(await status_async(details=args.details), args.json)
    elif args.command == "self-test":
        result = await self_test_async()
        print_result(result, args.json)
        return 0 if result["ok"] else 2
    elif args.command == "add":
        content = read_content(args)
        if not content.strip():
            raise SystemExit("content is required")
        tags = [t.strip() for t in args.tags.split(",") if t.strip()]
        metadata = parse_json_arg(args.metadata, {})
        from agent_identity import build_identity
        identity_input = {"metadata": metadata, "model_id": args.model_id, "model_provider": args.provider, "harness": args.harness, "thread_id": args.thread_id, "agent_id": args.agent_id, "reasoning_effort": args.reasoning_effort}
        metadata["identity"] = build_identity(identity_input, agent=args.agent, source=args.source, session_id=args.session_id)
        print_result(
            await add_event_async(
                agent=args.agent,
                event_type=args.event_type,
                space=args.space,
                memory_kind=args.memory_kind,
                role=args.role,
                source=args.source,
                session_id=args.session_id,
                summary=args.summary,
                tags=tags,
                metadata=metadata,
                content=content,
            ),
            args.json,
        )
    elif args.command == "search":
        tags = [t.strip() for t in args.tags.split(",") if t.strip()]
        print_result(
            await search_events_async(
                args.query,
                limit=args.limit,
                space=args.space,
                all_spaces=args.all_spaces,
                agent=args.agent,
                event_type=args.event_type,
                memory_kind=args.memory_kind,
                tags=tags,
                min_score=args.min_score,
            ),
            args.json,
        )
    elif args.command == "grep":
        tags = [t.strip() for t in args.tags.split(",") if t.strip()]
        print_result(
            await grep_events_async(
                args.query,
                limit=args.limit,
                space=args.space,
                all_spaces=args.all_spaces,
                agent=args.agent,
                event_type=args.event_type,
                memory_kind=args.memory_kind,
                tags=tags,
                since=args.since,
                contains=args.contains,
                timeout_ms=args.timeout_ms,
            ),
            args.json,
        )
    elif args.command == "audit":
        print_result(await audit_events_async(args.limit, args.session_id, args.agent, args.space, args.all_spaces, model_id=args.model_id, harness=args.harness, thread_id=args.thread_id, provider=args.provider), args.json)
    elif args.command == "trace":
        rows = await trace_archives_async(
            archive_id=args.archive_id,
            event_id=args.event_id,
            event_uid=args.event_uid,
            payload_sha256=args.sha256,
            session_id=args.session_id,
            space=args.space,
            all_spaces=args.all_spaces,
            limit=args.limit,
            include_payload=not args.no_payload,
        )
        if args.content_only:
            print("\n\n".join(str(row.get("payload", "")) for row in rows))
        else:
            print_result(rows, args.json)
    elif args.command == "backfill-legacy-archives":
        print_result(await backfill_legacy_archives_async(limit=args.limit, dry_run=args.dry_run), args.json)
    elif args.command == "backfill-identities":
        from identity_history import enrich_identities
        conn = await connect()
        try:
            print_result(await enrich_identities(conn, limit=args.limit, before_id=args.before_id, since=parse_datetime_arg(args.since), apply=args.apply), args.json)
        finally:
            await conn.close()
    elif args.command == "ingest-codex-transcript":
        from codex_transcript_import import import_transcript
        conn = await connect()
        try:
            result = await import_transcript(conn, path=args.path, limit=args.limit, apply=args.apply)
            print_result(result, args.json)
            return 0 if result["ok"] else 2
        finally:
            await conn.close()
    elif args.command == "ingest-error-log":
        print_result(
            await ingest_error_log_async(
                log_file=args.log_file,
                target_space=args.target_space,
                limit_groups=args.limit_groups,
                dry_run=args.dry_run,
            ),
            args.json,
        )
    elif args.command == "ingest-file":
        tags = [t.strip() for t in args.tags.split(",") if t.strip()]
        print_result(
            await ingest_file_async(
                args.path,
                space=args.space,
                agent=args.agent,
                source=args.source,
                session_id=args.session_id,
                tags=tags,
                chunk_chars=args.chunk_chars,
                chunk_overlap=args.chunk_overlap,
                archive=not args.no_archive,
                dry_run=args.dry_run,
            ),
            args.json,
        )
    elif args.command == "ingest-path":
        tags = [t.strip() for t in args.tags.split(",") if t.strip()]
        print_result(
            await ingest_path_async(
                args.path,
                space=args.space,
                agent=args.agent,
                source=args.source,
                session_id=args.session_id,
                tags=tags,
                chunk_chars=args.chunk_chars,
                chunk_overlap=args.chunk_overlap,
                archive=not args.no_archive,
                dry_run=args.dry_run,
                recursive=args.recursive,
                extensions=args.extensions,
                max_files=args.max_files,
            ),
            args.json,
        )
    elif args.command == "spaces":
        metadata = parse_json_arg(args.metadata, {})
        if args.space:
            space = await ensure_space_async(
                args.space,
                description=args.description,
                parent_space=args.parent_space,
                metadata=metadata,
            )
            print_result({"ok": True, "space": space}, args.json)
        else:
            print_result(await list_spaces_async(), args.json)
    elif args.command == "entities":
        print_result(
            await search_entities_async(args.query, space=args.space, all_spaces=args.all_spaces, limit=args.limit),
            args.json,
        )
    elif args.command == "relate":
        print_result(
            await relate_entities_async(
                source=args.source,
                target=args.target,
                relation=args.relation,
                space=args.space,
                description=args.description,
                metadata=parse_json_arg(args.metadata, {}),
            ),
            args.json,
        )
    elif args.command == "graph":
        print_result(
            await graph_neighbors_async(args.entity, space=args.space, all_spaces=args.all_spaces, limit=args.limit),
            args.json,
        )
    return 0


def main() -> int:
    args = build_parser().parse_args()
    try:
        return asyncio.run(run_cli(args))
    except BrokenPipeError:
        return 1
    except Exception as exc:
        write_error_log(exc)
        if args.json:
            print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        else:
            print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
